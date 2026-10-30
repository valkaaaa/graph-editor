"""Дерево объектов, редактор таблицы, редактор матрицы."""
import os
import uuid
import tkinter as tk

from tkinter import ttk, messagebox, Menu, filedialog, simpledialog
from translator import tr


class TreeUIMixin:

    def _get_selected_indices(self):
        result = []
        for item in self.graph_tree.selection():
            vals = self.graph_tree.item(item, 'values')
            if vals:
                try:
                    result.append(int(vals[0]))
                except (ValueError, IndexError):
                    pass
        return result

    def _on_tree_double_click(self, event):
        item = self.graph_tree.identify_row(event.y)
        if not item:
            return
        vals = self.graph_tree.item(item, 'values')
        if not vals:
            self._toggle_folder_collapse(item)
            return
        self.graph_tree.selection_set(item)
        self._load_selected_to_form()
        self._refresh_stats()
        self.notebook.select(0)

    def _toggle_folder_collapse(self, item):
        text = self.graph_tree.item(item, 'text')
        name = text.replace("📁 ", "").replace("📂 ", "")
        if name in self.hidden_folders:
            self.hidden_folders.discard(name)
        else:
            self.hidden_folders.add(name)
        self._update_listbox()

    def _create_group(self):
        """Объединяет выбранные объекты в визуальную группу."""
        sel = self._get_selected_indices()
        if len(sel) < 2:
            return

        group_id = str(uuid.uuid4())[:8]
        group_name = simpledialog.askstring(
            "Группа", "Название группы:", parent=self.root
        ) or f"Группа {group_id}"

        self._snapshot("Создана группа")
        for idx in sel:
            self.graphs_data[idx]['group_id'] = group_id
            self.graphs_data[idx]['group_name'] = group_name

        self._update_listbox()
        self._redraw()

    def _on_tree_right_click(self, event):
        item = self.graph_tree.identify_row(event.y)
        if not item:
            return
        self._refresh_stats()
        if item not in self.graph_tree.selection():
            self.graph_tree.selection_set(item)
        selection = self._get_selected_indices()
        vals = self.graph_tree.item(item, 'values')
        menu = Menu(self.root, tearoff=0)

        if len(selection) > 1:
            menu.add_command(label=tr.t('ctx.group_color', count=len(selection)),
                             command=self._group_change_color)
            menu.add_command(label=tr.t('ctx.group_lw', count=len(selection)),
                             command=self._group_change_lw)
            menu.add_command(label=tr.t('ctx.group_ms', count=len(selection)),
                             command=self._group_change_ms)
            menu.add_command(label=tr.t('ctx.group_palette'),
                             command=self._show_palette_dialog)
            menu.add_command(label=tr.t('ctx.group_toggle_visibility', count=len(selection)),
                             command=self._toggle_visibility)

            # --- НОВОЕ: Группировка ---
            menu.add_separator()
            menu.add_command(label="📦 Сгруппировать в одну папку",
                             command=self._create_group)

            menu.add_separator()
            menu.add_command(label=tr.t('ctx.group_move', count=len(selection)),
                             command=self._move_to_folder)
            menu.add_command(label=tr.t('ctx.group_export', count=len(selection)),
                             command=self._export_selected_json)
            menu.add_separator()
            menu.add_command(label=tr.t('ctx.group_duplicate', count=len(selection)),
                             command=self._duplicate_graph)
            menu.add_command(label=tr.t('ctx.group_delete', count=len(selection)),
                             command=self._delete_graph)
            menu.post(event.x_root, event.y_root)
            return

        if vals:
            menu.add_command(label=tr.t('ctx.edit'), command=self._load_selected_to_form)
            menu.add_command(label=tr.t('ctx.duplicate'), command=self._duplicate_graph)

            # --- НОВОЕ: Слои (Z-order) ---
            menu.add_separator()
            menu.add_command(label="На передний план",
                             command=lambda: self._move_zorder('front'))
            menu.add_command(label="На задний план",
                             command=lambda: self._move_zorder('back'))
            menu.add_command(label="Поднять на слой выше",
                             command=lambda: self._move_zorder('up'))
            menu.add_command(label="Опустить на слой ниже",
                             command=lambda: self._move_zorder('down'))

            menu.add_separator()
            menu.add_command(label=tr.t('ctx.notes'), command=self._edit_notes)
            menu.add_command(label=tr.t('ctx.rename'), command=self._rename_graph)
            menu.add_command(label=tr.t('ctx.move_folder'), command=self._move_to_folder)
            menu.add_command(label=tr.t('ctx.toggle_visibility'), command=self._toggle_visibility)
            menu.add_separator()
            menu.add_command(label=tr.t('ctx.delete'), command=self._delete_graph)
        else:
            name = (self.graph_tree.item(item, 'text').replace("📁 ", "").replace("📂 ", ""))
            menu.add_command(label=tr.t('ctx.rename_folder'), command=lambda: self._rename_folder(name))
            menu.add_command(label=tr.t('ctx.toggle_folder'), command=lambda: self._toggle_folder_collapse(item))
            menu.add_separator()
            menu.add_command(label=tr.t('ctx.export_folder_csv'), command=lambda: self._export_folder_csv(name))
            menu.add_command(label=tr.t('ctx.export_folder_json'), command=lambda: self._export_folder_json(name))
            menu.add_separator()
            menu.add_command(label=tr.t('ctx.delete_folder'), command=lambda: self._delete_folder(item, False))
            menu.add_command(label=tr.t('ctx.delete_folder_content'), command=lambda: self._delete_folder(item, True))
        menu.post(event.x_root, event.y_root)

    def _open_table_editor(self, axis=None):
        t = tr.find_type_by_display(self.type_var.get())

        matrix_types = (
            "Heatmap", "Contour", "Imshow", "3D Поверхность", "3D Сетка",
            "Chord", "Adjacency matrix", "Дендрограмма", "Cluster dendrogram",
            "Alluvial diagram", "Stream plot", "Streamlines", "Mapper algorithm",
            "Bump chart", "Parallel coordinates", "Radar chart", "Isopleth",
            "Isosurface 3D", "Waterfall spectrogram 3D", "Choropleth",
            "Dot density map", "Диаграмма Тейлора", "Полярная теплокарта"
        )

        if t in matrix_types:
            self._open_matrix_editor()
            return

        win = tk.Toplevel(self.root)
        win.title(tr.t('table.editor_title'))
        win.geometry("780x600")
        win.transient(self.root)
        skip_x = t in ("Круговая", "Кольцевая", "Вложенная круговая",
                       "Мозаика", "Гистограмма", "Treemap", "Sunburst")
        skip_z = ("3D" not in t and t not in ("Quiver", "Пузырьковая",
                                              "3D Полигон", "3D Триангуляция"))
        skip_yerr = (t not in ("Линия", "Точечная", "Столбчатая") and t != "Quiver")

        cols = []
        if not skip_x:
            cols.append(('x', self.entry_x, 'X'))
        cols.append(('y', self.entry_y, 'Y'))
        if not skip_z:
            cols.append(('z', self.entry_z, 'Z'))
        if not skip_yerr:
            cols.append(('yerr', self.entry_yerr, 'Yerr'))

        series = {}
        for cid, entry, _ in cols:
            text = entry.get().strip()
            series[cid] = self._parse_data(text) if text else []
            if series[cid] is None:
                series[cid] = []

        max_len = max((len(v) for v in series.values()), default=0)
        if max_len == 0:
            max_len = 5

        for cid in series:
            while len(series[cid]) < max_len:
                series[cid].append('')

        top = ttk.Frame(win, padding=5)
        top.pack(fill=tk.X)
        rows_var = tk.StringVar(value=tr.t('table.rows', n=max_len))
        ttk.Label(top, textvariable=rows_var).pack(side=tk.LEFT)

        tree_frame = ttk.Frame(win)
        tree_frame.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        col_ids = ['idx'] + [c[0] for c in cols]

        tree = ttk.Treeview(tree_frame, columns=col_ids, show='headings',
                            height=18, selectmode='extended')
        for cid in col_ids:
            text = '#' if cid == 'idx' else cid.upper()
            tree.heading(cid, text=text)
            tree.column(cid, width=40 if cid == 'idx' else 110, anchor=tk.CENTER)

        tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        vsb = ttk.Scrollbar(tree_frame, orient=tk.VERTICAL, command=tree.yview)
        vsb.pack(side=tk.RIGHT, fill=tk.Y)
        tree.configure(yscrollcommand=vsb.set)

        def refresh_indices():
            for i, item in enumerate(tree.get_children(), 1):
                vals = list(tree.item(item, 'values'))
                vals[0] = i
                tree.item(item, values=vals)

        def update_rows_label():
            rows_var.set(tr.t('table.rows', n=len(tree.get_children())))

        def add_row(values=None):
            if values is None:
                values = [''] * len(cols)
            tree.insert('', tk.END, values=[0] + list(values))
            refresh_indices()
            update_rows_label()

        for i in range(max_len):
            row = [series[c[0]][i] if i < len(series[c[0]]) else '' for c in cols]
            add_row(row)

        def on_double_click(event):
            region = tree.identify('region', event.x, event.y)
            if region != 'cell':
                return
            item = tree.identify_row(event.y)
            col = tree.identify_column(event.x)
            if not item or col == '#1':
                return
            bbox = tree.bbox(item, col)
            if not bbox:
                return
            ee = ttk.Entry(win)
            ee.place(x=bbox[0], y=bbox[1], width=bbox[2], height=bbox[3])
            ee.insert(0, tree.set(item, col))
            ee.select_range(0, tk.END)
            ee.focus()

            def commit(e=None):
                val = ee.get().strip()
                if val:
                    try:
                        float(val.replace(',', '.'))
                    except ValueError:
                        ee.destroy()
                        return
                tree.set(item, col, val)
                ee.destroy()

            ee.bind('<Return>', commit)
            ee.bind('<Escape>', lambda e: ee.destroy())
            ee.bind('<FocusOut>', commit)

        tree.bind('<Double-1>', on_double_click)

        row_btns = ttk.Frame(win, padding=5)
        row_btns.pack(fill=tk.X)

        def delete_selected():
            for item in tree.selection():
                tree.delete(item)
            refresh_indices()
            update_rows_label()

        def clear_all():
            if messagebox.askyesno(tr.t('msg.confirm'), 'Очистить таблицу?'):
                for item in tree.get_children():
                    tree.delete(item)
                update_rows_label()

        def sort_by_col(col_id):
            items = []
            for item in tree.get_children():
                val = tree.set(item, col_id)
                try:
                    key = float(val.replace(',', '.')) if val else float('inf')
                except (ValueError, TypeError):
                    key = float('inf')
                items.append((key, item))
            items.sort()
            for _, item in items:
                tree.move(item, '', 'end')
            refresh_indices()

        ttk.Button(row_btns, text=tr.t('table.btn_add_row'),
                   command=add_row).pack(side=tk.LEFT, padx=2)
        ttk.Button(row_btns, text=tr.t('table.btn_add_10'),
                   command=lambda: [add_row() for _ in range(10)]).pack(side=tk.LEFT, padx=2)
        ttk.Button(row_btns, text=tr.t('table.btn_remove_row'),
                   command=delete_selected).pack(side=tk.LEFT, padx=2)
        ttk.Button(row_btns, text=tr.t('table.btn_clear'),
                   command=clear_all).pack(side=tk.LEFT, padx=2)
        if any(c[0] == 'x' for c in cols):
            ttk.Button(row_btns, text=tr.t('table.btn_sort_x'),
                       command=lambda: sort_by_col('x')).pack(side=tk.LEFT, padx=2)

        io_btns = ttk.Frame(win, padding=5)
        io_btns.pack(fill=tk.X)

        def paste_from_clipboard():
            try:
                cb = win.clipboard_get()
            except tk.TclError:
                return
            for item in tree.get_children():
                tree.delete(item)
            for line in cb.replace('\r\n', '\n').strip().split('\n'):
                parts = [p.strip() for p in line.split('\t')]
                row_vals = [''] * len(cols)
                for i, p in enumerate(parts[:len(cols)]):
                    row_vals[i] = p
                add_row(row_vals)
            update_rows_label()

        def export_excel_table():
            try:
                import openpyxl
            except ImportError:
                messagebox.showerror(tr.t('msg.error'), tr.t('excel.install_openpyxl'))
                return
            path = filedialog.asksaveasfilename(
                defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
            if not path:
                return
            try:
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Data"
                for i, (_, _, hdr) in enumerate(cols, 1):
                    ws.cell(row=1, column=i, value=hdr)
                for r, item in enumerate(tree.get_children(), 2):
                    for i, (cid, _, _) in enumerate(cols, 1):
                        val = tree.set(item, cid).strip()
                        try:
                            ws.cell(row=r, column=i,
                                    value=float(val.replace(',', '.')) if val else None)
                        except (ValueError, TypeError):
                            ws.cell(row=r, column=i, value=val or None)
                wb.save(path)
                self._show_history(tr.t('excel.exported', name=os.path.basename(path)))
            except Exception as e:
                messagebox.showerror(tr.t('msg.error'), str(e))

        def import_excel_table():
            try:
                import openpyxl
            except ImportError:
                messagebox.showerror(tr.t('msg.error'), tr.t('excel.install_openpyxl'))
                return
            path = filedialog.askopenfilename(
                filetypes=[("Excel", "*.xlsx"), ("All", "*.*")])
            if not path:
                return
            try:
                wb = openpyxl.load_workbook(path, data_only=True)
            except Exception as e:
                messagebox.showerror(tr.t('excel.bad_file'), str(e))
                return
            sheet_win = tk.Toplevel(win)
            sheet_win.title(tr.t('excel.sheet_choose'))
            sheet_win.transient(win)
            sheet_win.geometry("320x240")
            ttk.Label(sheet_win, text=tr.t('excel.sheet_choose')).pack(padx=10, pady=(10, 4))
            lb = tk.Listbox(sheet_win, height=7)
            lb.pack(fill=tk.BOTH, expand=True, padx=10)
            for name in wb.sheetnames:
                lb.insert(tk.END, name)
            lb.selection_set(0)

            def do_it():
                sel = lb.curselection()
                if not sel:
                    return
                sheet_name = wb.sheetnames[sel[0]]
                sheet_win.destroy()
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    return
                first = [str(c).strip().lower() if c is not None else '' for c in rows[0]]
                has_header = any(h in ('x', 'y', 'z', 'yerr', 'err') for h in first)
                data_rows = rows[1:] if has_header else rows
                for item in tree.get_children():
                    tree.delete(item)
                for row in data_rows:
                    if all(c is None or str(c).strip() == '' for c in row):
                        continue
                    vals = ['' if c is None else c for c in row[:len(cols)]]
                    while len(vals) < len(cols):
                        vals.append('')
                    add_row(vals)
                update_rows_label()

            ttk.Button(sheet_win, text=tr.t('excel.apply'), command=do_it).pack(pady=10)

        ttk.Button(io_btns, text=tr.t('table.btn_paste'),
                   command=paste_from_clipboard).pack(side=tk.LEFT, padx=2)
        ttk.Button(io_btns, text=tr.t('table.btn_import_excel'),
                   command=import_excel_table).pack(side=tk.LEFT, padx=2)
        ttk.Button(io_btns, text=tr.t('table.btn_export_excel'),
                   command=export_excel_table).pack(side=tk.LEFT, padx=2)

        bottom = ttk.Frame(win, padding=5)
        bottom.pack(fill=tk.X)

        def apply_and_close():
            for cid, entry, _ in cols:
                values = []
                for item in tree.get_children():
                    val = tree.set(item, cid).strip()
                    if val:
                        values.append(val)
                entry.delete(0, tk.END)
                entry.insert(0, ", ".join(values))
            win.destroy()
            self._on_data_change()

        ttk.Button(bottom, text=tr.t('table.btn_cancel'),
                   command=win.destroy).pack(side=tk.RIGHT, padx=2)
        ttk.Button(bottom, text=tr.t('table.btn_apply'),
                   command=apply_and_close).pack(side=tk.RIGHT, padx=2)
        win.bind('<Control-v>', lambda e: paste_from_clipboard() or 'break')
        win.bind('<Delete>', lambda e: delete_selected())
        win.bind('<Escape>', lambda e: win.destroy())

    def _open_matrix_editor(self):
        win = tk.Toplevel(self.root)
        win.title(tr.t('table.matrix_title'))
        win.geometry("600x500")
        win.transient(self.root)
        ttk.Label(win, text=tr.t('table.matrix_hint'),
                  foreground='gray').pack(anchor=tk.W, padx=10, pady=(10, 2))
        txt = tk.Text(win, height=20, font=('Consolas', 10), undo=True)
        txt.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        current = self.entry_z.get().strip()
        if current:
            for r in current.split(';'):
                txt.insert(tk.END, r.strip() + '\n')

        def paste_clipboard():
            try:
                cb = win.clipboard_get()
            except tk.TclError:
                return
            txt.delete('1.0', tk.END)
            txt.insert('1.0', cb)
            return 'break'

        def apply_matrix():
            content = txt.get('1.0', tk.END).strip()
            lines = [l.strip() for l in content.split('\n') if l.strip()]
            self.entry_z.delete(0, tk.END)
            self.entry_z.insert(0, ' ; '.join(lines))
            win.destroy()
            self._on_data_change()

        bf = ttk.Frame(win, padding=5)
        bf.pack(fill=tk.X)
        ttk.Button(bf, text=tr.t('table.btn_paste'),
                   command=paste_clipboard).pack(side=tk.LEFT, padx=2)
        ttk.Button(bf, text=tr.t('table.btn_cancel'),
                   command=win.destroy).pack(side=tk.RIGHT, padx=2)
        ttk.Button(bf, text=tr.t('table.btn_apply'),
                   command=apply_matrix).pack(side=tk.RIGHT, padx=2)
        win.bind('<Control-v>', lambda e: paste_clipboard())
        win.bind('<Escape>', lambda e: win.destroy())