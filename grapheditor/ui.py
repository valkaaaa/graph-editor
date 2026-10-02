import os
import re
import sys
import tkinter as tk
import numpy as np
import matplotlib as mpl

from tkinter import (ttk, messagebox, colorchooser,
                     filedialog, Menu, simpledialog)

from translator import (tr, INTERNAL_TYPES, INTERNAL_ANIM_TYPES,
                        SUPPORTED_LANGS)
from settings import FONTS, UNIT_TO_M, TOOLTIPS, attach_tooltip
from themes import (THEMES, THEME_KEYS, is_custom_theme,
                    save_custom_theme, delete_custom_theme,
                    theme_key_from_name)


EMOJI_FALLBACK = {
    'toolbar.open_json':   '📂',
    'toolbar.save_json':   '💾',
    'toolbar.zip':         '🗜️',
    'toolbar.import_csv':  '⬇️',
    'toolbar.export_csv':  '⬆️',
    'toolbar.export_html': '🌐',
    'toolbar.save_image':  '🖼️',
    'toolbar.undo':        '↶',
    'toolbar.redo':        '↷',
    'create.create_btn':   '➕',
    'create.apply_btn':    '✓',
    'create.palette':      '🎨',
    'gen.generate_btn':    '⚙️',
    'gen.random_btn':      '🎲',
    'obj.load_btn':        '📥',
    'obj.duplicate_btn':   '📋',
    'obj.delete_btn':      '🗑️',
    'obj.clear_btn':       '🧹',
    'obj.add_trend':       '📈',
    'right.anim_start':    '▶',
    'right.anim_pause':    '⏸',
    'right.anim_stop':     '■',
    'right.anim_save':     '💾',
    'right.anim_save_file':'📦',
    'right.anim_load_file':'📂',
    'stats.fit_linear':    '📏',
    'stats.fit_log':       '📐',
    'stats.fit_exp':       '📉',
    'stats.fit_power':     '📊',
    'stats.find_peaks':    '🔍',
    'right.apply':         '✓',
}


class UIMixin:

    # ========================================================================
    # ИКОНКИ И РЕСУРСЫ
    # ========================================================================
    def _assets_dir(self):
        if getattr(sys, 'frozen', False):
            base = getattr(sys, '_MEIPASS',
                           os.path.dirname(sys.executable))
        else:
            base = os.path.dirname(os.path.abspath(__file__))
        return os.path.join(base, 'assets')

    def _set_window_icon(self):
        if sys.platform == 'win32':
            try:
                import ctypes
                ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(
                    "GraphEditorPro.1.0.0")
            except Exception:
                pass
        ico = os.path.join(self._assets_dir(), 'icon.ico')
        png = os.path.join(self._assets_dir(), 'icon.png')
        try:
            if os.name == 'nt' and os.path.exists(ico):
                self.root.iconbitmap(ico)
            elif os.path.exists(png):
                self._window_icon = tk.PhotoImage(file=png)
                self.root.iconphoto(True, self._window_icon)
        except Exception as e:
            print(f"[icon] {e}")

    def _load_icon(self, name, size=20):
        if not hasattr(self, '_icon_cache'):
            self._icon_cache = {}
        cache_key = (name, size)
        if cache_key in self._icon_cache:
            return self._icon_cache[cache_key]
        path = os.path.join(self._assets_dir(), 'icons', f'{name}.png')
        if not os.path.exists(path):
            return None
        try:
            from PIL import Image, ImageTk
            img = Image.open(path).convert('RGBA')
            img = img.resize((size, size), Image.LANCZOS)
            photo = ImageTk.PhotoImage(img)
            self._icon_cache[cache_key] = photo
            return photo
        except Exception:
            return None

    # ========================================================================
    # i18n
    # ========================================================================
    def _tr_label(self, parent, key, **kw):
        w = ttk.Label(parent, text=tr.t(key))
        self._i18n_widgets.append((w, 'text', key, {}))
        for k, v in kw.items():
            try:
                w.config(**{k: v})
            except tk.TclError:
                pass
        return w

    def _tr_button(self, parent, key, command, **kw):
        emoji = EMOJI_FALLBACK.get(key, '')
        text = f"{emoji} {tr.t(key)}".strip() if emoji else tr.t(key)
        w = ttk.Button(parent, text=text, command=command, **kw)
        self._i18n_widgets.append((w, 'text', key, {}))
        tip = TOOLTIPS.get(key)
        if tip:
            attach_tooltip(w, tip)
        return w

    def _tr_button_icon(self, parent, key, icon_name, command,
                        icon_size=20, **kw):
        photo = self._load_icon(icon_name, icon_size)
        text = tr.t(key)
        if photo is None:
            return self._tr_button(parent, key, command, **kw)
        w = ttk.Button(parent, text=text, image=photo, compound='left',
                       command=command, **kw)
        self._i18n_widgets.append((w, 'text', key, {}))
        tip = TOOLTIPS.get(key)
        if tip:
            attach_tooltip(w, tip)
        return w

    def _tr_checkbutton(self, parent, key, variable, command=None, **kw):
        w = ttk.Checkbutton(parent, text=tr.t(key),
                            variable=variable, command=command, **kw)
        self._i18n_widgets.append((w, 'text', key, {}))
        return w

    def _tr_labelframe(self, parent, key, **kw):
        w = ttk.LabelFrame(parent, text=tr.t(key), **kw)
        self._i18n_widgets.append((w, 'text', key, {}))
        return w

    def _switch_lang(self, lang):
        if lang == tr.lang:
            return
        tr.set_lang(lang)
        self.settings['language'] = lang
        new_lang_name = tr.t('settings.lang_' + lang)

        for widget, attr, key, kwargs in list(self._i18n_widgets):
            try:
                has_image = False
                try:
                    if isinstance(widget, ttk.Button):
                        has_image = bool(widget.cget('image'))
                except Exception:
                    pass
                if has_image:
                    widget.config(text=tr.t(key, **kwargs))
                else:
                    emoji = EMOJI_FALLBACK.get(key, '')
                    text = (f"{emoji} {tr.t(key, **kwargs)}".strip()
                            if emoji else tr.t(key, **kwargs))
                    widget.config(text=text)
            except tk.TclError:
                self._i18n_widgets.remove((widget, attr, key, kwargs))

        self._build_menu()

        self.type_combo['values'] = [tr.type_name(t)
                                     for t in INTERNAL_TYPES]
        current_internal = tr.find_type_by_display(self.type_var.get())
        self.type_var.set(tr.type_name(current_internal))

        self.anim_combo['values'] = [tr.anim_name(a)
                                     for a in INTERNAL_ANIM_TYPES]
        current_anim = tr.find_anim_by_display(self.anim_type_var.get())
        self.anim_type_var.set(tr.anim_name(current_anim))

        self.notebook.tab(0, text=tr.t('tab.create'))
        self.notebook.tab(1, text=tr.t('tab.generator'))
        self.notebook.tab(2, text=tr.t('tab.objects'))
        self.notebook.tab(3, text=tr.t('tab.stats'))

        self.rnd_shape_combo['values'] = [
            tr.t('gen.rnd_shape_2d'),
            tr.t('gen.rnd_shape_3d'),
            tr.t('gen.rnd_shape_3d_grid')]

        self.status_var.set(tr.t('status.ready'))
        self._update_listbox()
        self._redraw()
        self._show_history(tr.t('hist.lang_changed', lang=new_lang_name))

    # ========================================================================
    # МЕНЮ
    # ========================================================================
    def _build_menu(self):
        for widget in self.root.winfo_children():
            if isinstance(widget, Menu):
                widget.destroy()

        self.menubar = Menu(self.root)
        self.root.config(menu=self.menubar)

        # Файл
        self.filemenu = Menu(self.menubar, tearoff=0)
        self.menubar.add_cascade(label=tr.t('menu.file'), menu=self.filemenu)
        self.filemenu.add_command(label=tr.t('menu.file.open_json'),
                                  accelerator="Ctrl+O", command=self._load_json)
        self.filemenu.add_command(label=tr.t('menu.file.save_json'),
                                  accelerator="Ctrl+S", command=self._save_json)
        self.filemenu.add_separator()
        self.filemenu.add_command(label=tr.t('menu.file.open_zip'), command=self._load_zip)
        self.filemenu.add_command(label=tr.t('menu.file.save_zip'), command=self._save_zip)
        self.filemenu.add_separator()
        self.filemenu.add_command(label=tr.t('file.open_graphx'),
                                  command=self._load_graphx)
        self.filemenu.add_command(label=tr.t('file.save_graphx'),
                                  command=self._save_graphx)
        self.filemenu.add_separator()
        self.filemenu.add_command(label=tr.t('menu.file.import_csv'), command=self._import_csv)
        self.filemenu.add_command(label=tr.t('menu.file.export_csv'), command=self._export_csv)
        self.filemenu.add_command(label=tr.t('menu.file.import_excel'), command=self._import_excel)
        self.filemenu.add_command(label=tr.t('menu.file.export_excel'), command=self._export_excel)
        self.filemenu.add_separator()
        self.filemenu.add_command(label=tr.t('menu.file.save_image'), command=self._save_image)
        self.filemenu.add_command(label=tr.t('menu.file.export_html'), command=self._export_html)
        self.filemenu.add_separator()
        self.filemenu.add_command(label=tr.t('menu.file.console'),
                                  accelerator="Ctrl+`", command=self._open_console)
        self.filemenu.add_command(label=tr.t('menu.file.tour'), command=self._start_onboarding)
        self.filemenu.add_separator()
        self.filemenu.add_command(label=tr.t('menu.file.settings'), command=self._open_settings)

        # Правка
        self.editmenu = Menu(self.menubar, tearoff=0)
        self.menubar.add_cascade(label=tr.t('menu.edit'), menu=self.editmenu)
        self.editmenu.add_command(label=tr.t('menu.edit.undo'),
                                  accelerator="Ctrl+Z", command=self._undo)
        self.editmenu.add_command(label=tr.t('menu.edit.redo'),
                                  accelerator="Ctrl+Y", command=self._redo)
        self.editmenu.add_separator()
        self.editmenu.add_command(label=tr.t('menu.edit.duplicate'),
                                  accelerator="Ctrl+D", command=self._duplicate_graph)
        self.editmenu.add_command(label=tr.t('menu.edit.notes'), command=self._edit_notes)
        self.editmenu.add_separator()
        self.editmenu.add_command(label=tr.t('menu.edit.clear_all'), command=self._clear_all)

        # Шаблоны
        self.templatemenu = Menu(self.menubar, tearoff=0)
        self.menubar.add_cascade(label=tr.t('menu.templates'), menu=self.templatemenu)
        for key, name in [
            ('menu.templates.empty', 'empty'),
            ('menu.templates.sine', 'sine'),
            ('menu.templates.damped', 'damped'),
            ('menu.templates.distribution', 'distribution'),
            ('menu.templates.spiral3d', 'spiral3d'),
        ]:
            self.templatemenu.add_command(
                label=tr.t(key), command=lambda n=name: self._apply_template(n))
        self.templatemenu.add_separator()
        for key, name in [
            ('menu.templates.trig', 'trig'),
            ('menu.templates.distributions', 'distributions'),
            ('menu.templates.dashboard', 'dashboard'),
            ('menu.templates.fitting', 'fitting'),
            ('menu.templates.showcase', 'showcase'),
            ('menu.templates.fractal', 'fractal'),
        ]:
            self.templatemenu.add_command(
                label=tr.t(key), command=lambda n=name: self._apply_template(n))
        self.templatemenu.add_separator()
        self.templatemenu.add_command(label=tr.t('menu.templates.save_current'),
                                      command=self._save_as_template)
        self.templatemenu.add_command(label=tr.t('menu.templates.open_folder'),
                                      command=self._open_templates_folder)
        tpls = self._list_user_templates()
        if tpls:
            sub = Menu(self.templatemenu, tearoff=0)
            for name, path in tpls:
                sub.add_command(label=name,
                                command=lambda p=path: self._load_template_file(p))
            self.templatemenu.add_cascade(label=tr.t('menu.templates.my'), menu=sub)
        self.templatemenu.add_separator()
        self.templatemenu.add_command(label=tr.t('menu.templates.export'),
                                      command=self._export_template_file)
        self.templatemenu.add_command(label=tr.t('menu.templates.import'),
                                      command=self._import_template_file)

        # Вид
        self.viewmenu = Menu(self.menubar, tearoff=0)
        self.menubar.add_cascade(label=tr.t('menu.view'), menu=self.viewmenu)
        for key, var, cmd in [
            ('menu.view.grid', self.show_grid_var, self._update_grid),
            ('menu.view.legend', self.show_legend_var, self._update_legend),
            ('menu.view.axes', self.show_axes_var, self._update_axes),
        ]:
            self.viewmenu.add_checkbutton(label=tr.t(key), variable=var, command=cmd)
        self.viewmenu.add_separator()
        self.viewmenu.add_checkbutton(label=tr.t('menu.view.smooth'),
                                      variable=self.smooth_var, command=self._redraw_all)
        self.viewmenu.add_checkbutton(label=tr.t('menu.view.tooltip'),
                                      variable=self.show_tooltip_var, command=self._redraw_all)
        self.viewmenu.add_checkbutton(label=tr.t('menu.view.crosshair'),
                                      variable=self.crosshair_enabled,
                                      command=self._toggle_crosshair)
        self.viewmenu.add_checkbutton(label=tr.t('menu.view.drag_points'),
                                      variable=self.drag_points_enabled)
        self.viewmenu.add_separator()
        self.viewmenu.add_checkbutton(label=tr.t('menu.view.panel_settings'),
                                      variable=self.show_settings_var,
                                      command=self._toggle_ui_panels)
        self.viewmenu.add_checkbutton(label=tr.t('menu.view.panel_limits'),
                                      variable=self.show_limits_var,
                                      command=self._toggle_ui_panels)
        self.viewmenu.add_checkbutton(label=tr.t('menu.view.panel_anim'),
                                      variable=self.show_anim_var,
                                      command=self._toggle_ui_panels)
        self.viewmenu.add_separator()

        # Подменю тем
        self.theme_menu = Menu(self.viewmenu, tearoff=0)
        self.viewmenu.add_cascade(label=tr.t('menu.view.theme'),
                                   menu=self.theme_menu)
        if not hasattr(self, 'theme_var'):
            self.theme_var = tk.StringVar(
                value=self.settings.get('theme', 'light'))
        self._populate_theme_menu()

        self.viewmenu.add_checkbutton(label=tr.t('menu.view.live_preview'),
                                      variable=self.live_preview_var)

    def _populate_theme_menu(self):
        """Заполняет подменю тем."""
        menu = self.theme_menu
        menu.delete(0, 'end')

        current = self.settings.get('theme', 'light')

        built_in = [(k, v) for k, v in THEMES.items()
                    if not v.get('custom', False)]
        for key, t in built_in:
            menu.add_radiobutton(
                label=t.get('name', key), variable=self.theme_var,
                value=key, command=lambda k=key: self._set_theme(k))

        custom = [(k, v) for k, v in THEMES.items()
                  if v.get('custom', False)]
        if custom:
            menu.add_separator()
            for key, t in custom:
                menu.add_radiobutton(
                    label='⭐ ' + t.get('name', key),
                    variable=self.theme_var, value=key,
                    command=lambda k=key: self._set_theme(k))

        menu.add_separator()
        menu.add_command(label=tr.t('menu.view.theme.create'),
                          command=lambda: self._open_theme_editor(None))
        if is_custom_theme(current):
            menu.add_command(label=tr.t('menu.view.theme.edit'),
                              command=lambda: self._open_theme_editor(current))
            menu.add_command(label=tr.t('menu.view.theme.delete'),
                              command=lambda: self._delete_theme(current))

    def _delete_theme(self, name):
        if not is_custom_theme(name):
            messagebox.showinfo(tr.t('msg.info'),
                                 tr.t('theme.editor.cannot_delete_builtin'))
            return
        display = THEMES.get(name, {}).get('name', name)
        if not messagebox.askyesno(tr.t('msg.confirm'),
                                     f"Удалить тему «{display}»?"):
            return
        try:
            delete_custom_theme(name)
            THEMES.pop(name, None)
        except Exception as e:
            messagebox.showerror(tr.t('msg.error'), str(e))
            return
        self._set_theme('light')
        self.theme_var.set('light')
        self._populate_theme_menu()
        self._show_history(tr.t('theme.deleted', name=display))

    def _open_theme_editor(self, name=None):
        """Диалог создания/редактирования темы."""
        # Редактирование встроенной — делаем копию
        if name and not is_custom_theme(name):
            base = dict(THEMES[name])
            base['name'] = tr.t('theme.copy_of',
                                 name=THEMES[name].get('name', name))
            name = None
        else:
            base = dict(THEMES.get(name, THEMES['light']))

        win = tk.Toplevel(self.root)
        win.title(tr.t('theme.editor.title'))
        win.transient(self.root)
        win.geometry("500x640")

        top = ttk.Frame(win, padding=8)
        top.pack(fill=tk.X)

        ttk.Label(top, text=tr.t('theme.editor.name')).grid(
            row=0, column=0, sticky=tk.W)
        name_var = tk.StringVar(value=base.get('name', ''))
        ttk.Entry(top, textvariable=name_var, width=30).grid(
            row=0, column=1, padx=5, sticky=tk.EW)

        dark_var = tk.BooleanVar(value=base.get('is_dark', True))
        ttk.Checkbutton(top, text=tr.t('theme.editor.dark'),
                        variable=dark_var).grid(
            row=1, column=0, columnspan=2, sticky=tk.W, pady=(6, 0))

        ttk.Label(top, text=tr.t('theme.editor.mpl_style')).grid(
            row=2, column=0, sticky=tk.W, pady=(6, 0))
        mpl_style_var = tk.StringVar(
            value=base.get('mpl_style', 'dark_background'))
        ttk.Combobox(top, textvariable=mpl_style_var,
                     values=['default', 'dark_background', 'ggplot',
                             'seaborn-v0_8', 'bmh'],
                     state='readonly', width=20).grid(
            row=2, column=1, padx=5, sticky=tk.W, pady=(6, 0))
        top.columnconfigure(1, weight=1)

        colors_frame = ttk.LabelFrame(win,
                                       text=tr.t('theme.editor.colors'),
                                       padding=8)
        colors_frame.pack(fill=tk.BOTH, expand=True, padx=8, pady=8)

        color_vars = {}
        for key, label_key in THEME_KEYS:
            row = ttk.Frame(colors_frame)
            row.pack(fill=tk.X, pady=2)
            ttk.Label(row, text=tr.t(label_key) + ":",
                      width=22, anchor=tk.W).pack(side=tk.LEFT)

            var = tk.StringVar(value=base.get(key, '#000000'))
            color_vars[key] = var

            prev = tk.Label(row, text='  ', bg=var.get(), width=4,
                             relief=tk.SOLID, borderwidth=1)
            prev.pack(side=tk.LEFT, padx=(2, 6))

            ttk.Entry(row, textvariable=var, width=12).pack(
                side=tk.LEFT, padx=(0, 4))

            def make_picker(v=var, p=prev):
                def pick():
                    c = colorchooser.askcolor(
                        title=tr.t('theme.editor.pick'),
                        initialcolor=v.get())
                    if c[1]:
                        v.set(c[1]); p.config(bg=c[1])
                return pick
            ttk.Button(row, text=tr.t('theme.editor.pick'),
                        command=make_picker()).pack(side=tk.LEFT)

            def make_trace(v=var, p=prev):
                def update(*_):
                    try: p.config(bg=v.get())
                    except Exception: pass
                return update
            var.trace_add('write', make_trace())

        btn_frame = ttk.Frame(win, padding=8)
        btn_frame.pack(fill=tk.X)

        def collect():
            return {
                'name': name_var.get().strip(),
                'is_dark': dark_var.get(),
                'mpl_style': mpl_style_var.get(),
                'custom': True,
                **{k: v.get() for k, v in color_vars.items()},
            }

        def do_preview():
            THEMES['__preview__'] = collect()
            self._apply_theme('__preview__')

        def do_save():
            t = collect()
            display_name = t['name']
            if not display_name:
                messagebox.showerror(tr.t('msg.error'),
                                      tr.t('theme.editor.need_name'))
                return
            key = theme_key_from_name(display_name)
            if not key:
                messagebox.showerror(tr.t('msg.error'),
                                      tr.t('theme.editor.bad_name'))
                return
            if key in THEMES and not THEMES[key].get('custom', False):
                messagebox.showerror(tr.t('msg.error'),
                                      tr.t('theme.editor.bad_name'))
                return
            if key in THEMES and THEMES[key].get('custom', False):
                if not messagebox.askyesno(
                        tr.t('msg.confirm'),
                        tr.t('theme.editor.name_exists')):
                    return
            try:
                THEMES.pop('__preview__', None)
                save_custom_theme(key, t)
                t['custom'] = True
                THEMES[key] = t
            except Exception as e:
                messagebox.showerror(tr.t('msg.error'), str(e))
                return
            self.settings['theme'] = key
            self.theme_var.set(key)
            self._apply_theme(key)
            self._populate_theme_menu()
            self._show_history(tr.t('theme.saved', name=display_name))
            win.destroy()

        def do_cancel():
            THEMES.pop('__preview__', None)
            self._apply_theme(self.settings.get('theme', 'light'))
            win.destroy()

        ttk.Button(btn_frame, text=tr.t('theme.editor.preview'),
                    command=do_preview).pack(side=tk.LEFT, padx=2)
        ttk.Button(btn_frame, text=tr.t('theme.editor.cancel'),
                    command=do_cancel).pack(side=tk.RIGHT, padx=2)
        ttk.Button(btn_frame, text=tr.t('theme.editor.save'),
                    command=do_save).pack(side=tk.RIGHT, padx=2)
        win.protocol("WM_DELETE_WINDOW", do_cancel)

    # ========================================================================
    # UI — сборка
    # ========================================================================
    def _build_ui(self):
        tb = ttk.Frame(self.root, padding=2)
        tb.pack(side=tk.TOP, fill=tk.X)

        for key, icon, cmd in [
            ('toolbar.open_json', 'open', self._load_json),
            ('toolbar.save_json', 'save', self._save_json),
            ('toolbar.zip', 'zip', self._save_zip),
        ]:
            self._tr_button_icon(tb, key, icon, cmd).pack(side=tk.LEFT, padx=2)
        ttk.Separator(tb, orient=tk.VERTICAL).pack(side=tk.LEFT, fill=tk.Y, padx=5)
        for key, icon, cmd in [
            ('toolbar.import_csv', 'import', self._import_csv),
            ('toolbar.export_csv', 'export', self._export_csv),
            ('toolbar.export_html', 'html', self._export_html),
        ]:
            self._tr_button_icon(tb, key, icon, cmd).pack(side=tk.LEFT, padx=2)
        ttk.Separator(tb, orient=tk.VERTICAL).pack(side=tk.LEFT, fill=tk.Y, padx=5)
        self._tr_button_icon(tb, 'toolbar.save_image', 'image',
                              self._save_image).pack(side=tk.LEFT, padx=2)
        ttk.Separator(tb, orient=tk.VERTICAL).pack(side=tk.LEFT, fill=tk.Y, padx=5)
        self._tr_button_icon(tb, 'toolbar.undo', 'undo',
                              self._undo).pack(side=tk.LEFT, padx=2)
        self._tr_button_icon(tb, 'toolbar.redo', 'redo',
                              self._redo).pack(side=tk.LEFT, padx=2)

        left_frame = ttk.Frame(self.root, padding=5, width=440)
        left_frame.pack(side=tk.LEFT, fill=tk.Y)
        left_frame.pack_propagate(False)

        self.notebook = ttk.Notebook(left_frame)
        self.notebook.pack(fill=tk.BOTH, expand=True)

        # Создание
        tab_create = ttk.Frame(self.notebook, padding=5)
        self.notebook.add(tab_create, text=tr.t('tab.create'))

        self._tr_label(tab_create, 'create.plot_type').pack(anchor=tk.W)
        self.type_combo = ttk.Combobox(
            tab_create, textvariable=self.type_var,
            values=[tr.type_name(t) for t in INTERNAL_TYPES],
            state="readonly")
        self.type_combo.pack(fill=tk.X, pady=(0, 5))
        self.type_combo.bind("<<ComboboxSelected>>", self._on_type_change)

        grid_row = ttk.Frame(tab_create)
        grid_row.pack(fill=tk.X, pady=(0, 5))
        self._tr_label(grid_row, 'create.subplot_grid').pack(side=tk.LEFT)
        self.subplot_grid_combo = ttk.Combobox(
            grid_row, textvariable=self.subplot_grid_var,
            values=['1x1', '1x2', '2x1', '2x2', '2x3', '3x2', '3x3'],
            state='readonly', width=6)
        self.subplot_grid_combo.pack(side=tk.LEFT, padx=5)
        self.subplot_grid_combo.bind('<<ComboboxSelected>>', self._on_grid_change)
        self._tr_label(grid_row, 'create.subplot_index').pack(side=tk.LEFT, padx=(15, 0))
        self.subplot_index_spin = ttk.Spinbox(
            grid_row, from_=0, to=0, textvariable=self.subplot_index_var, width=4)
        self.subplot_index_spin.pack(side=tk.LEFT, padx=5)

        self.frame_x = ttk.Frame(tab_create)
        self.frame_y = ttk.Frame(tab_create)
        self.frame_z = ttk.Frame(tab_create)
        self.frame_yerr = ttk.Frame(tab_create)
        self.frame_pie_labels = ttk.Frame(tab_create)
        self.frame_pie_colors = ttk.Frame(tab_create)
        self.frame_bins = ttk.Frame(tab_create)
        self.frame_treemap = ttk.Frame(tab_create)
        self.frame_y2_check = ttk.Frame(tab_create)
        self.frame_fill = ttk.Frame(tab_create)
        self.frame_glow = ttk.Frame(tab_create)
        self.frame_hint = ttk.Frame(tab_create)

        # Единый редактор таблицы для всех полей
        self.entry_x = self._create_entry_with_table(
            self.frame_x, 'create.x_data', "1, 2, 3, 4, 5", 'x')
        self.entry_y = self._create_entry_with_table(
            self.frame_y, 'create.y_data', "10, 20, 15, 25, 30", 'y')
        self.entry_z = self._create_entry_with_table(
            self.frame_z, 'create.z_data', "5, 10, 5, 15, 10", 'z')
        self.entry_yerr = self._create_entry_with_table(
            self.frame_yerr, 'create.yerr', "", 'yerr')

        self.entry_pie = self._create_entry(
            self.frame_pie_labels, 'create.pie_labels',
            "Яблоки, Груши, Сливы, Персики, Абрикосы")
        self.entry_pie_colors = self._create_entry(
            self.frame_pie_colors, 'create.pie_colors',
            "red, blue, green, orange, purple")
        self.treemap_parents = self._create_entry(
            self.frame_treemap, 'create.treemap_parents', "")

        self._tr_label(self.frame_bins, 'create.bins').pack(side=tk.LEFT)
        self.bins_var = tk.IntVar(value=10)
        ttk.Spinbox(self.frame_bins, from_=1, to=200,
                    textvariable=self.bins_var, width=6).pack(side=tk.LEFT, padx=5)

        self.hint_label = ttk.Label(self.frame_hint, text="", foreground="#666",
                                     font=('Arial', 8), wraplength=420,
                                     justify=tk.LEFT)
        self.hint_label.pack(anchor=tk.W)

        self.frame_x.pack(fill=tk.X)
        self.frame_y.pack(fill=tk.X)

        self.use_y2_var = tk.BooleanVar(value=False)
        self.fill_var = tk.BooleanVar(value=False)
        self.glow_var = tk.BooleanVar(value=False)

        self.frame_y2_check.pack(fill=tk.X)
        self._tr_checkbutton(self.frame_y2_check, 'create.use_y2',
                              variable=self.use_y2_var).pack(anchor=tk.W)
        self.frame_fill.pack(fill=tk.X)
        self._tr_checkbutton(self.frame_fill, 'create.fill',
                              variable=self.fill_var).pack(anchor=tk.W)
        self.frame_glow.pack(fill=tk.X)
        self._tr_checkbutton(self.frame_glow, 'create.glow',
                              variable=self.glow_var).pack(anchor=tk.W)

        # Единицы
        units_frame = self._tr_labelframe(tab_create, 'units.title', padding=5)
        units_frame.pack(fill=tk.X, pady=5)
        ur = ttk.Frame(units_frame)
        ur.pack(fill=tk.X)
        ALL_UNITS = ["—", "m", "cm", "mm", "km", "µm", "nm",
                     "in", "ft", "yd", "mi", "nmi",
                     "s", "ms", "µs", "min", "h",
                     "kg", "g", "mg", "t", "lb", "oz"]
        for i, (key, var_name) in enumerate([
            ('units.x_label', 'unit_x_var'),
            ('units.y_label', 'unit_y_var'),
            ('units.z_label', 'unit_z_var'),
        ]):
            self._tr_label(ur, key).grid(row=0, column=i * 2, sticky=tk.W,
                                           padx=(8 if i else 0, 0))
            v = tk.StringVar(value="—")
            setattr(self, var_name, v)
            ttk.Combobox(ur, textvariable=v, values=ALL_UNITS,
                         width=5, state="readonly").grid(
                row=0, column=i * 2 + 1, padx=3)

        # Внешний вид
        style_frame = self._tr_labelframe(tab_create, 'create.appearance', padding=5)
        style_frame.pack(fill=tk.X, pady=5)
        cr = ttk.Frame(style_frame)
        cr.pack(fill=tk.X, pady=2)
        self._tr_label(cr, 'create.color').pack(side=tk.LEFT)
        self.color_var = tk.StringVar(value="blue")
        ttk.Entry(cr, textvariable=self.color_var, width=10).pack(side=tk.LEFT, padx=5)
        self._tr_button_icon(cr, 'create.palette', 'palette',
                              self._pick_color).pack(side=tk.LEFT, expand=True, fill=tk.X)

        gf = ttk.Frame(style_frame)
        gf.pack(fill=tk.X, pady=2)
        self._tr_label(gf, 'create.linewidth').grid(row=0, column=0, sticky=tk.W)
        self.lw_var = tk.DoubleVar(value=2.0)
        ttk.Spinbox(gf, from_=0.5, to=10, increment=0.5,
                    textvariable=self.lw_var, width=5).grid(row=0, column=1, padx=5)
        self._tr_label(gf, 'create.linestyle').grid(row=1, column=0, sticky=tk.W)
        self.ls_var = tk.StringVar(value="-")
        ttk.Combobox(gf, textvariable=self.ls_var,
                     values=["-", "--", ":", "-."], width=5,
                     state="readonly").grid(row=1, column=1, padx=5)
        self._tr_label(gf, 'create.markersize').grid(row=2, column=0, sticky=tk.W)
        self.ms_var = tk.DoubleVar(value=6.0)
        ttk.Spinbox(gf, from_=1, to=30, increment=1,
                    textvariable=self.ms_var, width=5).grid(row=2, column=1, padx=5)
        self._tr_label(gf, 'create.marker').grid(row=3, column=0, sticky=tk.W)
        self.marker_var = tk.StringVar(value="o")
        ttk.Combobox(gf, textvariable=self.marker_var,
                     values=["None", "o", "s", "^", "v", "D", "*", "x", "+"],
                     width=5, state="readonly").grid(row=3, column=1, padx=5)

        self._tr_label(tab_create, 'create.name').pack(anchor=tk.W, pady=(5, 0))
        self.entry_label = ttk.Entry(tab_create)
        self.entry_label.pack(fill=tk.X, pady=(0, 5))
        self._tr_label(tab_create, 'create.notes').pack(anchor=tk.W)
        self.entry_notes = ttk.Entry(tab_create)
        self.entry_notes.pack(fill=tk.X, pady=(0, 5))

        bf = ttk.Frame(tab_create)
        bf.pack(fill=tk.X, pady=5)
        self._tr_button_icon(bf, 'create.create_btn', 'create',
                              self._add_graph).pack(
            side=tk.LEFT, expand=True, fill=tk.X, padx=2)
        self._tr_button_icon(bf, 'create.apply_btn', 'apply',
                              self._update_graph).pack(
            side=tk.RIGHT, expand=True, fill=tk.X, padx=2)

        # Генератор
        tab_gen = ttk.Frame(self.notebook, padding=5)
        self.notebook.add(tab_gen, text=tr.t('tab.generator'))
        self._tr_label(tab_gen, 'gen.formula_label',
                        font=('Arial', 9, 'bold')).pack(anchor=tk.W, pady=(0, 2))
        self._tr_label(tab_gen, 'gen.formula_hint',
                        foreground="gray", wraplength=400).pack(anchor=tk.W)
        self.gen_expr = ttk.Entry(tab_gen)
        self.gen_expr.pack(fill=tk.X, pady=5)
        self.gen_expr.insert(0, "sin(x) * exp(-x/5)")

        rf = ttk.Frame(tab_gen)
        rf.pack(fill=tk.X, pady=5)
        self._tr_label(rf, 'gen.from').pack(side=tk.LEFT)
        self.gen_xmin = ttk.Entry(rf, width=6)
        self.gen_xmin.pack(side=tk.LEFT, padx=2)
        self.gen_xmin.insert(0, "0")
        self._tr_label(rf, 'gen.to').pack(side=tk.LEFT)
        self.gen_xmax = ttk.Entry(rf, width=6)
        self.gen_xmax.pack(side=tk.LEFT, padx=2)
        self.gen_xmax.insert(0, "10")
        self._tr_label(rf, 'gen.points').pack(side=tk.LEFT, padx=(10, 0))
        self.gen_n = ttk.Entry(rf, width=5)
        self.gen_n.pack(side=tk.LEFT, padx=2)
        self.gen_n.insert(0, "50")
        self._tr_label(rf, 'gen.unit').pack(side=tk.LEFT, padx=(10, 2))
        self.gen_unit_var = tk.StringVar(value="m")
        self.gen_unit_combo = ttk.Combobox(
            rf, textvariable=self.gen_unit_var,
            values=list(UNIT_TO_M.keys()), width=5, state="readonly")
        self.gen_unit_combo.pack(side=tk.LEFT)

        self._tr_button_icon(tab_gen, 'gen.generate_btn', 'generate',
                              self._generate_data).pack(fill=tk.X, pady=10)

        # Случайные
        rnd_frame = self._tr_labelframe(tab_gen, 'gen.random_title', padding=5)
        rnd_frame.pack(fill=tk.X, pady=5)
        r1 = ttk.Frame(rnd_frame)
        r1.pack(fill=tk.X, pady=2)
        self._tr_label(r1, 'gen.distribution').pack(side=tk.LEFT)
        self.rnd_dist_var = tk.StringVar(value="normal")
        ttk.Combobox(r1, textvariable=self.rnd_dist_var,
                     values=["normal", "uniform", "exponential",
                             "poisson", "integers"],
                     width=12, state="readonly").pack(side=tk.LEFT, padx=5)
        self._tr_label(r1, 'gen.n').pack(side=tk.LEFT, padx=(10, 2))
        self.rnd_n_var = tk.IntVar(value=50)
        ttk.Spinbox(r1, from_=1, to=100000,
                    textvariable=self.rnd_n_var, width=7).pack(side=tk.LEFT)

        r_shape = ttk.Frame(rnd_frame)
        r_shape.pack(fill=tk.X, pady=2)
        self._tr_label(r_shape, 'gen.rnd_shape').pack(side=tk.LEFT)
        self.rnd_shape_var = tk.StringVar(value=tr.t('gen.rnd_shape_2d'))
        self.rnd_shape_combo = ttk.Combobox(
            r_shape, textvariable=self.rnd_shape_var,
            values=[tr.t('gen.rnd_shape_2d'), tr.t('gen.rnd_shape_3d'),
                    tr.t('gen.rnd_shape_3d_grid')],
            state='readonly', width=26)
        self.rnd_shape_combo.pack(side=tk.LEFT, padx=5)
        self.rnd_shape_combo.bind('<<ComboboxSelected>>',
                                    self._on_rnd_shape_change)

        self.rnd_grid_frame = ttk.Frame(rnd_frame)
        self._tr_label(self.rnd_grid_frame, 'gen.rnd_grid_size').pack(side=tk.LEFT)
        self.rnd_grid_n_var = tk.IntVar(value=15)
        ttk.Spinbox(self.rnd_grid_frame, from_=3, to=200,
                    textvariable=self.rnd_grid_n_var, width=6).pack(side=tk.LEFT, padx=3)
        ttk.Label(self.rnd_grid_frame, text="×").pack(side=tk.LEFT)
        self.rnd_grid_m_var = tk.IntVar(value=15)
        ttk.Spinbox(self.rnd_grid_frame, from_=3, to=200,
                    textvariable=self.rnd_grid_m_var, width=6).pack(side=tk.LEFT, padx=3)

        r2 = ttk.Frame(rnd_frame)
        r2.pack(fill=tk.X, pady=2)
        self._tr_label(r2, 'gen.param1').pack(side=tk.LEFT)
        self.rnd_p1_var = tk.StringVar(value="0")
        ttk.Entry(r2, textvariable=self.rnd_p1_var, width=8).pack(side=tk.LEFT, padx=3)
        self._tr_label(r2, 'gen.param2').pack(side=tk.LEFT, padx=(8, 0))
        self.rnd_p2_var = tk.StringVar(value="1")
        ttk.Entry(r2, textvariable=self.rnd_p2_var, width=8).pack(side=tk.LEFT, padx=3)
        r3 = ttk.Frame(rnd_frame)
        r3.pack(fill=tk.X, pady=2)
        self._tr_label(r3, 'gen.seed').pack(side=tk.LEFT)
        self.rnd_seed_var = tk.StringVar(value="")
        ttk.Entry(r3, textvariable=self.rnd_seed_var, width=8).pack(side=tk.LEFT, padx=3)
        self._tr_button_icon(rnd_frame, 'gen.random_btn', 'dice',
                              self._generate_random_data).pack(fill=tk.X, pady=(5, 0))

        # Объекты
        tab_list = ttk.Frame(self.notebook, padding=5)
        self.notebook.add(tab_list, text=tr.t('tab.objects'))
        tree_frame = ttk.Frame(tab_list)
        tree_frame.pack(fill=tk.BOTH, expand=True)
        self.graph_tree = ttk.Treeview(tree_frame, show='tree', height=12,
                                        selectmode='extended')
        self.graph_tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, pady=(0, 5))
        tree_sb = ttk.Scrollbar(tree_frame, orient=tk.VERTICAL,
                                 command=self.graph_tree.yview)
        tree_sb.pack(side=tk.RIGHT, fill=tk.Y)
        self.graph_tree.configure(yscrollcommand=tree_sb.set)
        self.graph_tree.bind('<Double-1>', self._on_tree_double_click)
        self.graph_tree.bind('<Button-3>', self._on_tree_right_click)

        trend_frame = self._tr_labelframe(tab_list, 'obj.analysis', padding=5)
        trend_frame.pack(fill=tk.X, pady=(0, 5))
        trend_row = ttk.Frame(trend_frame)
        trend_row.pack(fill=tk.X)
        self._tr_label(trend_row, 'obj.trend_degree').pack(side=tk.LEFT)
        self.trend_degree_var = tk.IntVar(value=1)
        ttk.Combobox(trend_row, textvariable=self.trend_degree_var,
                     values=[1, 2, 3], width=3, state="readonly").pack(side=tk.LEFT, padx=5)
        self._tr_button_icon(trend_row, 'obj.add_trend', 'trend',
                              self._add_trendline).pack(side=tk.RIGHT, fill=tk.X, expand=True)

        bl = ttk.Frame(tab_list)
        bl.pack(fill=tk.X)
        for key, icon, cmd in [
            ('obj.load_btn', 'load', self._load_selected_to_form),
            ('obj.duplicate_btn', 'duplicate', self._duplicate_graph),
            ('obj.delete_btn', 'delete', self._delete_graph),
        ]:
            self._tr_button_icon(bl, key, icon, cmd).pack(
                side=tk.LEFT, expand=True, fill=tk.X, padx=2)
        self._tr_button_icon(bl, 'obj.clear_btn', 'clear',
                              self._clear_all).pack(
            side=tk.RIGHT, expand=True, fill=tk.X, padx=2)

        # Статистика
        tab_stats = ttk.Frame(self.notebook, padding=5)
        self.notebook.add(tab_stats, text=tr.t('tab.stats'))
        self.stats_label = ttk.Label(
            tab_stats, text=tr.t('stats.none_selected'),
            font=('Arial', 10, 'bold'), foreground='#666')
        self.stats_label.pack(anchor=tk.W, pady=(0, 5))

        self.stats_frame = self._tr_labelframe(tab_stats, 'stats.title', padding=5)
        self.stats_frame.pack(fill=tk.X, pady=5)
        self.stats_vars = {}
        for label_key, var_name in [
            ('stats.n', 'n'), ('stats.min', 'min'), ('stats.max', 'max'),
            ('stats.mean', 'mean'), ('stats.median', 'median'),
            ('stats.std', 'std'), ('stats.var', 'var'),
            ('stats.q1', 'q1'), ('stats.q3', 'q3'), ('stats.iqr', 'iqr'),
            ('stats.sum', 'sum'), ('stats.corr', 'corr'),
        ]:
            row = ttk.Frame(self.stats_frame)
            row.pack(fill=tk.X, pady=1)
            self._tr_label(row, label_key).pack(side=tk.LEFT)
            v = tk.StringVar(value='—')
            self.stats_vars[var_name] = v
            ttk.Label(row, textvariable=v,
                      font=('Consolas', 10)).pack(side=tk.RIGHT)

        fit_frame = self._tr_labelframe(tab_stats, 'stats.fit_title', padding=5)
        fit_frame.pack(fill=tk.X, pady=5)
        fr1 = ttk.Frame(fit_frame); fr1.pack(fill=tk.X, pady=2)
        self._tr_button_icon(fr1, 'stats.fit_linear', 'fit_linear',
                              lambda: self._add_fit('linear')).pack(side=tk.LEFT, expand=True, fill=tk.X, padx=1)
        self._tr_button_icon(fr1, 'stats.fit_log', 'fit_log',
                              lambda: self._add_fit('log')).pack(side=tk.LEFT, expand=True, fill=tk.X, padx=1)
        fr2 = ttk.Frame(fit_frame); fr2.pack(fill=tk.X, pady=2)
        self._tr_button_icon(fr2, 'stats.fit_exp', 'fit_exp',
                              lambda: self._add_fit('exp')).pack(side=tk.LEFT, expand=True, fill=tk.X, padx=1)
        self._tr_button_icon(fr2, 'stats.fit_power', 'fit_power',
                              lambda: self._add_fit('power')).pack(side=tk.LEFT, expand=True, fill=tk.X, padx=1)
        fr3 = ttk.Frame(fit_frame); fr3.pack(fill=tk.X, pady=2)
        self._tr_label(fr3, 'stats.fit_poly').pack(side=tk.LEFT)
        self.fit_degree_var = tk.IntVar(value=2)
        ttk.Spinbox(fr3, from_=1, to=10,
                    textvariable=self.fit_degree_var, width=4).pack(side=tk.LEFT, padx=5)
        self._tr_button_icon(fr3, 'obj.add_trend', 'trend',
                              self._add_fit_poly).pack(side=tk.RIGHT, fill=tk.X, expand=True)

        self.fit_result_var = tk.StringVar(value='')
        ttk.Label(tab_stats, textvariable=self.fit_result_var,
                  font=('Consolas', 9), foreground='#333',
                  wraplength=400, justify=tk.LEFT).pack(anchor=tk.W, pady=5, fill=tk.X)

        peaks_frame = self._tr_labelframe(tab_stats, 'stats.peaks_title', padding=5)
        peaks_frame.pack(fill=tk.X, pady=5)
        self._tr_button_icon(peaks_frame, 'stats.find_peaks', 'search',
                              self._find_peaks).pack(fill=tk.X)

        # Правая панель
        right_frame = ttk.Frame(self.root, padding=5)
        right_frame.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True)

        self.settings_frame = self._tr_labelframe(right_frame, 'right.canvas_settings', padding=5)
        self.settings_frame.pack(fill=tk.X, pady=(0, 5))
        for i, (key, var) in enumerate([
            ('right.title', self.plot_title_var),
            ('right.xlabel', self.plot_xlabel_var),
            ('right.ylabel', self.plot_ylabel_var),
        ]):
            self._tr_label(self.settings_frame, key).grid(row=i, column=0, sticky=tk.W)
            ttk.Entry(self.settings_frame, textvariable=var,
                      width=20).grid(row=i, column=1, padx=5, pady=2)
        self._tr_button_icon(self.settings_frame, 'right.apply', 'apply',
                              self._apply_canvas_settings).grid(
            row=0, column=2, rowspan=3, padx=5)

        self.limits_frame = self._tr_labelframe(right_frame, 'right.limits', padding=5)
        self.limits_frame.pack(fill=tk.X, pady=(0, 5))
        self.lim_auto_var = tk.BooleanVar(value=True)
        self._tr_checkbutton(self.limits_frame, 'right.limits_auto',
                              variable=self.lim_auto_var,
                              command=self._apply_limits).grid(
            row=0, column=0, rowspan=2, padx=5, sticky=tk.W)
        limits_conf = [
            ('right.limits_xmin', 'lim_xmin', 0, 1),
            ('right.limits_xmax', 'lim_xmax', 0, 3),
            ('right.limits_ymin', 'lim_ymin', 1, 1),
            ('right.limits_ymax', 'lim_ymax', 1, 3),
        ]
        for key, var_name, row, col in limits_conf:
            self._tr_label(self.limits_frame, key).grid(row=row, column=col,
                                                          sticky=tk.E, padx=2)
            v = tk.StringVar(value="")
            setattr(self, var_name, v)
            ttk.Entry(self.limits_frame, textvariable=v, width=8).grid(
                row=row, column=col + 1, padx=2, pady=2)
        self._tr_button_icon(self.limits_frame, 'right.apply', 'apply',
                              self._apply_limits).grid(
            row=0, column=5, rowspan=2, padx=10)

        self.anim_labelframe = self._tr_labelframe(right_frame, 'right.anim_engine', padding=5)
        self.anim_labelframe.pack(fill=tk.X, pady=(0, 5))
        ar = ttk.Frame(self.anim_labelframe); ar.pack(fill=tk.X, pady=2)
        self._tr_label(ar, 'right.anim_type').pack(side=tk.LEFT, padx=(0, 5))
        self.anim_combo = ttk.Combobox(
            ar, textvariable=self.anim_type_var,
            values=[tr.anim_name(a) for a in INTERNAL_ANIM_TYPES],
            state="readonly", width=22)
        self.anim_combo.pack(side=tk.LEFT, padx=(0, 10))
        self.anim_combo.bind("<<ComboboxSelected>>", self._on_anim_type_change)
        self.anim_speed_var = tk.IntVar(value=30)
        self._tr_label(ar, 'right.anim_frame_ms').pack(side=tk.LEFT, padx=(10, 2))
        ttk.Spinbox(ar, from_=10, to=500, increment=10,
                    textvariable=self.anim_speed_var, width=5).pack(side=tk.LEFT)
        self.frame_custom_anim = ttk.Frame(self.anim_labelframe)
        ttk.Label(self.frame_custom_anim,
                  text="y(x,t)/z(x,y,t):").pack(side=tk.LEFT)
        self.entry_custom_anim = ttk.Entry(self.frame_custom_anim, width=25)
        self.entry_custom_anim.pack(side=tk.LEFT, padx=5)
        self.entry_custom_anim.insert(0, "sin(x + t) * cos(t/2)")
        baf = ttk.Frame(self.anim_labelframe); baf.pack(fill=tk.X)
        for key, icon, cmd in [
            ('right.anim_start', 'play', self._start_animation),
            ('right.anim_pause', 'pause', self._toggle_pause_anim),
            ('right.anim_stop', 'stop', self._stop_animation),
            ('right.anim_save', 'save', self._save_animation),
            ('right.anim_save_file', 'package', self._save_graphxanim),
            ('right.anim_load_file', 'open', self._load_graphxanim),
        ]:
            self._tr_button_icon(baf, key, icon, cmd).pack(side=tk.LEFT, padx=2)
        self.anim_progress = ttk.Progressbar(
            self.anim_labelframe, maximum=self.anim_frames_total,
            mode='determinate')
        self.anim_progress.pack(fill=tk.X, padx=2, pady=(3, 0))

        self.plot_notebook = ttk.Notebook(right_frame)
        self.plot_notebook.pack(fill=tk.BOTH, expand=True)
        self.plot_notebook.bind('<<NotebookTabChanged>>',
                                self._on_plot_tab_changed)
        self._build_documents_bar(right_frame)
        self.plot_frame = None

        statusbar = ttk.Frame(self.root, relief=tk.SUNKEN)
        statusbar.pack(side=tk.BOTTOM, fill=tk.X)
        ttk.Label(statusbar, textvariable=self.status_var, anchor=tk.W,
                  padding=(5, 2)).pack(side=tk.LEFT, fill=tk.X, expand=True)
        ttk.Separator(statusbar, orient=tk.VERTICAL).pack(
            side=tk.LEFT, fill=tk.Y, padx=5)
        ttk.Label(statusbar, textvariable=self.history_var, anchor=tk.W,
                  padding=(5, 2), foreground="#008000").pack(
            side=tk.LEFT, fill=tk.X, expand=True)
        console_btn = ttk.Button(statusbar, text="🐍 Python", width=12,
                                  command=self._open_console)
        console_btn.pack(side=tk.RIGHT, padx=5, pady=1)
        attach_tooltip(console_btn, "Открыть Python-консоль (Ctrl+`)")

    # ========================================================================
    # ИСТОРИЯ / ШРИФТЫ / ХОТКЕИ
    # ========================================================================
    def _show_history(self, text):
        self.history_var.set(f"🕐 {text}")
        if self._history_after_id:
            try:
                self.root.after_cancel(self._history_after_id)
            except Exception:
                pass
        self._history_after_id = self.root.after(
            4000, lambda: self.history_var.set(""))

    def _apply_font_settings(self):
        s = self.settings
        mpl.rcParams['font.family'] = s['font_family']
        mpl.rcParams['font.size'] = s['font_base']
        mpl.rcParams['axes.titlesize'] = s['font_title']
        mpl.rcParams['axes.labelsize'] = s['font_labels']
        mpl.rcParams['legend.fontsize'] = s['font_legend']
        mpl.rcParams['xtick.labelsize'] = s['font_ticks']
        mpl.rcParams['ytick.labelsize'] = s['font_ticks']
        mpl.rcParams['axes.linewidth'] = s['axis_lw']
        if hasattr(self, 'canvas'):
            self._redraw_all()

    def _bind_hotkeys(self):
        b = self.root.bind
        # Файлы
        b('<Control-s>', lambda e: self._save_json())
        b('<Control-Shift-S>', lambda e: self._save_json_as())
        b('<Control-o>', lambda e: self._load_json())
        b('<Control-p>', lambda e: self._save_image())
        b('<Control-e>', lambda e: self._export_excel())
        # Правка
        b('<Control-z>', lambda e: self._undo())
        b('<Control-y>', lambda e: self._redo())
        b('<Control-Shift-Z>', lambda e: self._redo())
        b('<Control-d>', lambda e: self._duplicate_graph())
        b('<Control-Shift-N>', lambda e: self._clear_all())
        # Документы
        b('<Control-n>', lambda e: self._new_document(switch=True))
        b('<Control-t>', lambda e: self._new_document(switch=True))
        b('<Control-w>', lambda e: self._close_document(
            self.current_doc_idx))
        b('<Control-Tab>', lambda e: self._next_document())
        b('<Control-Shift-Tab>', lambda e: self._prev_document())
        # Прочее
        b('<Control-grave>', lambda e: self._open_console())
        b('<F5>', lambda e: self._redraw())
        b('<F11>', lambda e: self._toggle_fullscreen())
        b('<Delete>', self._on_delete_key)

    def _build_documents_bar(self, parent):
        bar = ttk.Frame(parent)
        bar.pack(fill=tk.X, before=self.plot_notebook)
        ttk.Button(bar, text="＋ Новый документ",
                   command=lambda: self._new_document(switch=True)
                   ).pack(side=tk.LEFT, padx=2, pady=2)
        ttk.Button(bar, text="✕ Закрыть",
                   command=lambda: self._close_document(self.current_doc_idx)
                   ).pack(side=tk.LEFT, padx=2, pady=2)

    def _toggle_fullscreen(self):
        self._fullscreen = not getattr(self, '_fullscreen', False)
        self.root.attributes('-fullscreen', self._fullscreen)

    def _on_delete_key(self, event=None):
        if self.root.focus_get() == self.graph_tree:
            self._delete_graph()
            return "break"

    # ========================================================================
    # ДЕРЕВО
    # ========================================================================
    def _get_selected_indices(self):
        result = []
        for item in self.graph_tree.selection():
            vals = self.graph_tree.item(item, 'values')
            if vals:
                try:
                    result.append(int(vals[0]))
                except (ValueError, IndexError):
                    pass
        return result

    def _on_tree_double_click(self, event):
        item = self.graph_tree.identify_row(event.y)
        if not item:
            return
        vals = self.graph_tree.item(item, 'values')
        if not vals:
            self._toggle_folder_collapse(item)
            return
        self.graph_tree.selection_set(item)
        self._load_selected_to_form()
        self._refresh_stats()
        self.notebook.select(0)

    def _toggle_folder_collapse(self, item):
        text = self.graph_tree.item(item, 'text')
        name = text.replace("📁 ", "").replace("📂 ", "")
        if name in self.hidden_folders:
            self.hidden_folders.discard(name)
        else:
            self.hidden_folders.add(name)
        self._update_listbox()

    def _on_tree_right_click(self, event):
        item = self.graph_tree.identify_row(event.y)
        if not item:
            return
        self._refresh_stats()
        if item not in self.graph_tree.selection():
            self.graph_tree.selection_set(item)
        selection = self._get_selected_indices()
        vals = self.graph_tree.item(item, 'values')
        menu = Menu(self.root, tearoff=0)

        if len(selection) > 1:
            menu.add_command(label=tr.t('ctx.group_color', count=len(selection)),
                              command=self._group_change_color)
            menu.add_command(label=tr.t('ctx.group_lw', count=len(selection)),
                              command=self._group_change_lw)
            menu.add_command(label=tr.t('ctx.group_ms', count=len(selection)),
                              command=self._group_change_ms)
            menu.add_separator()
            menu.add_command(label=tr.t('ctx.group_move', count=len(selection)),
                              command=self._move_to_folder)
            menu.add_command(label=tr.t('ctx.group_export', count=len(selection)),
                              command=self._export_selected_json)
            menu.add_separator()
            menu.add_command(label=tr.t('ctx.group_duplicate', count=len(selection)),
                              command=self._duplicate_graph)
            menu.add_command(label=tr.t('ctx.group_delete', count=len(selection)),
                              command=self._delete_graph)
            menu.post(event.x_root, event.y_root)
            return

        if vals:
            menu.add_command(label=tr.t('ctx.edit'), command=self._load_selected_to_form)
            menu.add_command(label=tr.t('ctx.duplicate'), command=self._duplicate_graph)
            menu.add_separator()
            menu.add_command(label=tr.t('ctx.notes'), command=self._edit_notes)
            menu.add_command(label=tr.t('ctx.rename'), command=self._rename_graph)
            menu.add_command(label=tr.t('ctx.move_folder'), command=self._move_to_folder)
            menu.add_separator()
            menu.add_command(label=tr.t('ctx.delete'), command=self._delete_graph)
        else:
            name = (self.graph_tree.item(item, 'text')
                    .replace("📁 ", "").replace("📂 ", ""))
            menu.add_command(label=tr.t('ctx.rename_folder'),
                              command=lambda: self._rename_folder(name))
            menu.add_command(label=tr.t('ctx.toggle_folder'),
                              command=lambda: self._toggle_folder_collapse(item))
            menu.add_separator()
            menu.add_command(label=tr.t('ctx.export_folder_csv'),
                              command=lambda: self._export_folder_csv(name))
            menu.add_command(label=tr.t('ctx.export_folder_json'),
                              command=lambda: self._export_folder_json(name))
            menu.add_separator()
            menu.add_command(label=tr.t('ctx.delete_folder'),
                              command=lambda: self._delete_folder(item, False))
            menu.add_command(label=tr.t('ctx.delete_folder_content'),
                              command=lambda: self._delete_folder(item, True))
        menu.post(event.x_root, event.y_root)

    # ========================================================================
    # СОЗДАНИЕ ПОЛЕЙ
    # ========================================================================
    def _create_entry(self, parent, label_key, default_text):
        self._tr_label(parent, label_key).pack(anchor=tk.W)
        entry = ttk.Entry(parent)
        entry.pack(fill=tk.X, pady=(0, 5))
        entry.insert(0, default_text)
        return entry

    def _create_entry_with_table(self, parent, label_key, default_text, axis):
        frame = ttk.Frame(parent)
        frame.pack(fill=tk.X, pady=(0, 5))
        self._tr_label(frame, label_key).pack(anchor=tk.W)
        entry = ttk.Entry(frame)
        entry.pack(side=tk.LEFT, fill=tk.X, expand=True)
        entry.insert(0, default_text)
        entry.bind("<Control-v>",
                   lambda e: self._paste_from_clipboard(e, entry))
        entry.bind("<KeyRelease>", self._on_data_change)
        self._tr_button(frame, 'create.table_btn',
                        lambda: self._open_table_editor(axis),
                        width=8).pack(side=tk.RIGHT, padx=(5, 0))
        return entry

    # ========================================================================
    # ТАБЛИЦА
    # ========================================================================
    def _open_table_editor(self, axis=None):
        """Улучшенный редактор таблицы: X/Y/Z/Yerr + Excel + буфер."""
        t = tr.find_type_by_display(self.type_var.get())
        if t in ("Heatmap", "Contour", "Imshow",
                 "3D Поверхность", "3D Сетка"):
            self._open_matrix_editor()
            return

        win = tk.Toplevel(self.root)
        win.title(tr.t('table.editor_title'))
        win.geometry("780x600")
        win.transient(self.root)

        skip_x = t in ("Круговая", "Кольцевая", "Вложенная круговая",
                        "Мозаика", "Гистограмма", "Treemap", "Sunburst")
        skip_z = ("3D" not in t
                  and t not in ("Quiver", "Пузырьковая",
                                "3D Полигон", "3D Триангуляция"))
        skip_yerr = (t not in ("Линия", "Точечная", "Столбчатая")
                     and t != "Quiver")

        cols = []
        if not skip_x:
            cols.append(('x', self.entry_x, 'X'))
        cols.append(('y', self.entry_y, 'Y'))
        if not skip_z:
            cols.append(('z', self.entry_z, 'Z'))
        if not skip_yerr:
            cols.append(('yerr', self.entry_yerr, 'Yerr'))

        series = {}
        for cid, entry, _ in cols:
            text = entry.get().strip()
            series[cid] = self._parse_data(text) if text else []
            if series[cid] is None:
                series[cid] = []
        max_len = max((len(v) for v in series.values()), default=0)
        if max_len == 0:
            max_len = 5
        for cid in series:
            while len(series[cid]) < max_len:
                series[cid].append('')

        top = ttk.Frame(win, padding=5); top.pack(fill=tk.X)
        rows_var = tk.StringVar(value=tr.t('table.rows', n=max_len))
        ttk.Label(top, textvariable=rows_var).pack(side=tk.LEFT)

        tree_frame = ttk.Frame(win)
        tree_frame.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        col_ids = ['idx'] + [c[0] for c in cols]
        tree = ttk.Treeview(tree_frame, columns=col_ids, show='headings',
                             height=18, selectmode='extended')
        for cid in col_ids:
            text = '#' if cid == 'idx' else cid.upper()
            tree.heading(cid, text=text)
            tree.column(cid, width=40 if cid == 'idx' else 110, anchor=tk.CENTER)
        tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        vsb = ttk.Scrollbar(tree_frame, orient=tk.VERTICAL, command=tree.yview)
        vsb.pack(side=tk.RIGHT, fill=tk.Y)
        tree.configure(yscrollcommand=vsb.set)

        def refresh_indices():
            for i, item in enumerate(tree.get_children(), 1):
                vals = list(tree.item(item, 'values'))
                vals[0] = i
                tree.item(item, values=vals)

        def update_rows_label():
            rows_var.set(tr.t('table.rows', n=len(tree.get_children())))

        def add_row(values=None):
            if values is None:
                values = [''] * len(cols)
            tree.insert('', tk.END, values=[0] + list(values))
            refresh_indices()
            update_rows_label()

        for i in range(max_len):
            row = [series[c[0]][i] if i < len(series[c[0]]) else ''
                   for c in cols]
            add_row(row)

        def on_double_click(event):
            region = tree.identify('region', event.x, event.y)
            if region != 'cell':
                return
            item = tree.identify_row(event.y)
            col = tree.identify_column(event.x)
            if not item or col == '#1':
                return
            bbox = tree.bbox(item, col)
            if not bbox:
                return
            ee = ttk.Entry(win)
            ee.place(x=bbox[0], y=bbox[1], width=bbox[2], height=bbox[3])
            ee.insert(0, tree.set(item, col))
            ee.select_range(0, tk.END)
            ee.focus()

            def commit(e=None):
                val = ee.get().strip()
                if val:
                    try:
                        float(val.replace(',', '.'))
                    except ValueError:
                        ee.destroy(); return
                tree.set(item, col, val)
                ee.destroy()

            ee.bind('<Return>', commit)
            ee.bind('<Escape>', lambda e: ee.destroy())
            ee.bind('<FocusOut>', commit)

        tree.bind('<Double-1>', on_double_click)

        row_btns = ttk.Frame(win, padding=5); row_btns.pack(fill=tk.X)

        def delete_selected():
            for item in tree.selection():
                tree.delete(item)
            refresh_indices(); update_rows_label()

        def clear_all():
            if messagebox.askyesno(tr.t('msg.confirm'), 'Очистить таблицу?'):
                for item in tree.get_children():
                    tree.delete(item)
                update_rows_label()

        def sort_by_col(col_id):
            items = []
            for item in tree.get_children():
                val = tree.set(item, col_id)
                try:
                    key = float(val.replace(',', '.')) if val else float('inf')
                except (ValueError, TypeError):
                    key = float('inf')
                items.append((key, item))
            items.sort()
            for _, item in items:
                tree.move(item, '', 'end')
            refresh_indices()

        ttk.Button(row_btns, text=tr.t('table.btn_add_row'),
                    command=add_row).pack(side=tk.LEFT, padx=2)
        ttk.Button(row_btns, text=tr.t('table.btn_add_10'),
                    command=lambda: [add_row() for _ in range(10)]).pack(side=tk.LEFT, padx=2)
        ttk.Button(row_btns, text=tr.t('table.btn_remove_row'),
                    command=delete_selected).pack(side=tk.LEFT, padx=2)
        ttk.Button(row_btns, text=tr.t('table.btn_clear'),
                    command=clear_all).pack(side=tk.LEFT, padx=2)
        if any(c[0] == 'x' for c in cols):
            ttk.Button(row_btns, text=tr.t('table.btn_sort_x'),
                        command=lambda: sort_by_col('x')).pack(side=tk.LEFT, padx=2)

        io_btns = ttk.Frame(win, padding=5); io_btns.pack(fill=tk.X)

        def paste_from_clipboard():
            try:
                cb = win.clipboard_get()
            except tk.TclError:
                return
            for item in tree.get_children():
                tree.delete(item)
            for line in cb.replace('\r\n', '\n').strip().split('\n'):
                parts = [p.strip() for p in line.split('\t')]
                row_vals = [''] * len(cols)
                for i, p in enumerate(parts[:len(cols)]):
                    row_vals[i] = p
                add_row(row_vals)
            update_rows_label()

        def export_excel_table():
            try:
                import openpyxl
            except ImportError:
                messagebox.showerror(tr.t('msg.error'),
                                      tr.t('excel.install_openpyxl'))
                return
            path = filedialog.asksaveasfilename(
                defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
            if not path:
                return
            try:
                wb = openpyxl.Workbook()
                ws = wb.active; ws.title = "Data"
                for i, (_, _, hdr) in enumerate(cols, 1):
                    ws.cell(row=1, column=i, value=hdr)
                for r, item in enumerate(tree.get_children(), 2):
                    for i, (cid, _, _) in enumerate(cols, 1):
                        val = tree.set(item, cid).strip()
                        try:
                            ws.cell(row=r, column=i,
                                     value=float(val.replace(',', '.')) if val else None)
                        except (ValueError, TypeError):
                            ws.cell(row=r, column=i, value=val or None)
                wb.save(path)
                self._show_history(tr.t('excel.exported',
                                          name=os.path.basename(path)))
            except Exception as e:
                messagebox.showerror(tr.t('msg.error'), str(e))

        def import_excel_table():
            try:
                import openpyxl
            except ImportError:
                messagebox.showerror(tr.t('msg.error'),
                                      tr.t('excel.install_openpyxl'))
                return
            path = filedialog.askopenfilename(
                filetypes=[("Excel", "*.xlsx"), ("All", "*.*")])
            if not path:
                return
            try:
                wb = openpyxl.load_workbook(path, data_only=True)
            except Exception as e:
                messagebox.showerror(tr.t('excel.bad_file'), str(e)); return

            sheet_win = tk.Toplevel(win)
            sheet_win.title(tr.t('excel.sheet_choose'))
            sheet_win.transient(win); sheet_win.geometry("320x240")
            ttk.Label(sheet_win, text=tr.t('excel.sheet_choose')).pack(padx=10, pady=(10, 4))
            lb = tk.Listbox(sheet_win, height=7)
            lb.pack(fill=tk.BOTH, expand=True, padx=10)
            for name in wb.sheetnames:
                lb.insert(tk.END, name)
            lb.selection_set(0)

            def do_it():
                sel = lb.curselection()
                if not sel:
                    return
                sheet_name = wb.sheetnames[sel[0]]
                sheet_win.destroy()
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    return
                first = [str(c).strip().lower() if c is not None else ''
                         for c in rows[0]]
                has_header = any(h in ('x', 'y', 'z', 'yerr', 'err')
                                  for h in first)
                data_rows = rows[1:] if has_header else rows
                for item in tree.get_children():
                    tree.delete(item)
                for row in data_rows:
                    if all(c is None or str(c).strip() == '' for c in row):
                        continue
                    vals = ['' if c is None else c for c in row[:len(cols)]]
                    while len(vals) < len(cols):
                        vals.append('')
                    add_row(vals)
                update_rows_label()

            ttk.Button(sheet_win, text=tr.t('excel.apply'),
                        command=do_it).pack(pady=10)

        ttk.Button(io_btns, text=tr.t('table.btn_paste'),
                    command=paste_from_clipboard).pack(side=tk.LEFT, padx=2)
        ttk.Button(io_btns, text=tr.t('table.btn_import_excel'),
                    command=import_excel_table).pack(side=tk.LEFT, padx=2)
        ttk.Button(io_btns, text=tr.t('table.btn_export_excel'),
                    command=export_excel_table).pack(side=tk.LEFT, padx=2)

        bottom = ttk.Frame(win, padding=5); bottom.pack(fill=tk.X)

        def apply_and_close():
            for cid, entry, _ in cols:
                values = []
                for item in tree.get_children():
                    val = tree.set(item, cid).strip()
                    if val:
                        values.append(val)
                entry.delete(0, tk.END)
                entry.insert(0, ", ".join(values))
            win.destroy()
            self._on_data_change()

        ttk.Button(bottom, text=tr.t('table.btn_cancel'),
                    command=win.destroy).pack(side=tk.RIGHT, padx=2)
        ttk.Button(bottom, text=tr.t('table.btn_apply'),
                    command=apply_and_close).pack(side=tk.RIGHT, padx=2)

        win.bind('<Control-v>', lambda e: paste_from_clipboard() or 'break')
        win.bind('<Delete>', lambda e: delete_selected())
        win.bind('<Escape>', lambda e: win.destroy())

    def _open_matrix_editor(self):
        """Редактор матрицы Z."""
        win = tk.Toplevel(self.root)
        win.title(tr.t('table.matrix_title'))
        win.geometry("600x500")
        win.transient(self.root)

        ttk.Label(win, text=tr.t('table.matrix_hint'),
                   foreground='gray').pack(anchor=tk.W, padx=10, pady=(10, 2))
        txt = tk.Text(win, height=20, font=('Consolas', 10), undo=True)
        txt.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)

        current = self.entry_z.get().strip()
        if current:
            for r in current.split(';'):
                txt.insert(tk.END, r.strip() + '\n')

        def paste_clipboard():
            try:
                cb = win.clipboard_get()
            except tk.TclError:
                return
            txt.delete('1.0', tk.END)
            txt.insert('1.0', cb)
            return 'break'

        def apply_matrix():
            content = txt.get('1.0', tk.END).strip()
            lines = [l.strip() for l in content.split('\n') if l.strip()]
            self.entry_z.delete(0, tk.END)
            self.entry_z.insert(0, ' ; '.join(lines))
            win.destroy()
            self._on_data_change()

        bf = ttk.Frame(win, padding=5); bf.pack(fill=tk.X)
        ttk.Button(bf, text=tr.t('table.btn_paste'),
                    command=paste_clipboard).pack(side=tk.LEFT, padx=2)
        ttk.Button(bf, text=tr.t('table.btn_cancel'),
                    command=win.destroy).pack(side=tk.RIGHT, padx=2)
        ttk.Button(bf, text=tr.t('table.btn_apply'),
                    command=apply_matrix).pack(side=tk.RIGHT, padx=2)
        win.bind('<Control-v>', lambda e: paste_clipboard())
        win.bind('<Escape>', lambda e: win.destroy())

    # ========================================================================
    # ПАНЕЛИ / НАСТРОЙКИ
    # ========================================================================
    def _toggle_ui_panels(self):
        anchor = getattr(self, 'plot_notebook', None) or self.plot_frame
        if anchor is None:
            return

        # --- Панель анимаций ---
        if self.show_anim_var.get():
            self.anim_labelframe.pack(fill=tk.X, pady=(0, 5), before=anchor)
        else:
            self.anim_labelframe.pack_forget()

        # --- Панель пределов ---
        if self.show_limits_var.get():
            target = self.anim_labelframe if self.show_anim_var.get() else anchor
            self.limits_frame.pack(fill=tk.X, pady=(0, 5), before=target)
        else:
            self.limits_frame.pack_forget()

        # --- Панель настроек холста ---
        if self.show_settings_var.get():
            if self.show_limits_var.get():
                target = self.limits_frame
            elif self.show_anim_var.get():
                target = self.anim_labelframe
            else:
                target = anchor
            self.settings_frame.pack(fill=tk.X, pady=(0, 5), before=target)
        else:
            self.settings_frame.pack_forget()

    def _open_settings(self):
        win = tk.Toplevel(self.root)
        win.title(tr.t('msg.settings_title'))
        win.transient(self.root)
        win.geometry("440x920")

        # Язык — выпадающий список
        ttk.Label(win, text=tr.t('settings.language'),
                  font=('Arial', 10, 'bold')).pack(anchor=tk.W, padx=10, pady=(10, 2))

        # Карта: отображаемая строка -> код языка
        label_to_code = {label: code for code, label in SUPPORTED_LANGS}
        code_to_label = {code: label for code, label in SUPPORTED_LANGS}

        current = self.settings.get('language', 'ru')
        if current not in code_to_label:
            current = 'en'

        lang_var = tk.StringVar(value=code_to_label[current])
        ttk.Combobox(win, textvariable=lang_var,
                     values=[lbl for _, lbl in SUPPORTED_LANGS],
                     state='readonly', width=30).pack(anchor=tk.W, padx=10)

        # Тема
        ttk.Label(win, text=tr.t('settings.theme'),
                  font=('Arial', 10, 'bold')).pack(anchor=tk.W, padx=10, pady=(15, 2))
        theme_var = tk.StringVar(value=self.settings.get('theme', 'light'))
        ttk.Combobox(win, textvariable=theme_var,
                     values=[k for k, _ in self._list_themes()],
                     state='readonly', width=25).pack(anchor=tk.W, padx=10)

        # Сглаживание
        ttk.Label(win, text=tr.t('settings.smooth'),
                  font=('Arial', 10, 'bold')).pack(anchor=tk.W, padx=10, pady=(15, 2))
        smooth_var = tk.IntVar(value=self.settings['smooth_window'])
        ttk.Label(win, text=tr.t('settings.smooth_window')).pack(anchor=tk.W, padx=10)
        ttk.Spinbox(win, from_=1, to=100, textvariable=smooth_var, width=10).pack(anchor=tk.W, padx=10)

        # Шрифты
        ttk.Label(win, text=tr.t('settings.fonts'),
                  font=('Arial', 10, 'bold')).pack(anchor=tk.W, padx=10, pady=(15, 2))
        font_family_var = tk.StringVar(value=self.settings['font_family'])
        ttk.Label(win, text=tr.t('settings.font_family')).pack(anchor=tk.W, padx=10)
        ttk.Combobox(win, textvariable=font_family_var, values=FONTS,
                     width=25, state="readonly").pack(anchor=tk.W, padx=10)
        font_vars = {}
        for key, label_key in [
            ('font_base', 'settings.font_base'), ('font_title', 'settings.font_title'),
            ('font_labels', 'settings.font_labels'), ('font_legend', 'settings.font_legend'),
            ('font_ticks', 'settings.font_ticks'),
        ]:
            ttk.Label(win, text=f"{tr.t(label_key)}:").pack(anchor=tk.W, padx=10, pady=(5, 0))
            var = tk.IntVar(value=self.settings[key])
            font_vars[key] = var
            ttk.Spinbox(win, from_=6, to=48, textvariable=var, width=8).pack(anchor=tk.W, padx=10)

        # Стиль
        ttk.Label(win, text=tr.t('settings.style'),
                  font=('Arial', 10, 'bold')).pack(anchor=tk.W, padx=10, pady=(15, 2))
        default_color_var = tk.StringVar(value=self.settings['default_color'])
        ttk.Label(win, text=tr.t('settings.default_color')).pack(anchor=tk.W, padx=10)
        ttk.Entry(win, textvariable=default_color_var, width=15).pack(anchor=tk.W, padx=10)
        default_lw_var = tk.DoubleVar(value=self.settings['default_lw'])
        ttk.Label(win, text=tr.t('settings.default_lw')).pack(anchor=tk.W, padx=10, pady=(5, 0))
        ttk.Spinbox(win, from_=0.5, to=10, increment=0.5,
                    textvariable=default_lw_var, width=8).pack(anchor=tk.W, padx=10)
        axis_lw_var = tk.DoubleVar(value=self.settings['axis_lw'])
        ttk.Label(win, text=tr.t('settings.axis_lw')).pack(anchor=tk.W, padx=10, pady=(5, 0))
        ttk.Spinbox(win, from_=0.5, to=5, increment=0.5,
                    textvariable=axis_lw_var, width=8).pack(anchor=tk.W, padx=10)

        # Обучение
        ttk.Label(win, text=tr.t('settings.tour'),
                  font=('Arial', 10, 'bold')).pack(anchor=tk.W, padx=10, pady=(15, 2))
        tour_var = tk.BooleanVar(value=self.settings.get('show_tour_on_start', False))
        ttk.Checkbutton(win, text=tr.t('settings.tour_show_on_start'),
                        variable=tour_var).pack(anchor=tk.W, padx=10)

        def run_tour_now():
            win.destroy()
            self.root.after(100, self._start_onboarding)

        ttk.Button(win, text=tr.t('settings.tour_run_now'),
                    command=run_tour_now).pack(anchor=tk.W, padx=10, pady=(5, 0))

        ttk.Label(win, text=tr.t('menu.view.drag_points'),
                  font=('Arial', 10, 'bold')).pack(anchor=tk.W, padx=10, pady=(15, 2))
        ttk.Checkbutton(win, text=tr.t('menu.view.drag_points'),
                        variable=self.drag_points_enabled).pack(anchor=tk.W, padx=10)

        ttk.Label(win, text=tr.t('settings.other'),
                  font=('Arial', 10, 'bold')).pack(anchor=tk.W, padx=10, pady=(15, 2))
        undo_var = tk.IntVar(value=self.max_undo)
        ttk.Label(win, text=tr.t('settings.max_undo')).pack(anchor=tk.W, padx=10)
        ttk.Spinbox(win, from_=10, to=500, textvariable=undo_var, width=10).pack(anchor=tk.W, padx=10)

        def save():
            # lang_var хранит отображаемую строку — превращаем в код
            new_lang = label_to_code.get(
                lang_var.get(),
                self.settings.get('language', 'ru'))

            lang_changed = new_lang != self.settings.get('language', 'ru')

            new_theme = theme_var.get()
            theme_changed = new_theme != self.settings.get('theme', 'light')

            self.settings['smooth_window'] = smooth_var.get()
            self.settings['font_family'] = font_family_var.get()
            self.settings['default_color'] = default_color_var.get()
            self.settings['default_lw'] = default_lw_var.get()
            self.settings['axis_lw'] = axis_lw_var.get()
            for key, var in font_vars.items():
                self.settings[key] = var.get()
            self.max_undo = undo_var.get()
            self.settings['show_tour_on_start'] = tour_var.get()
            self.settings['theme'] = new_theme
            self.settings['language'] = new_lang

            self.color_var.set(default_color_var.get())
            self.lw_var.set(default_lw_var.get())
            self._apply_font_settings()

            win.destroy()
            self._show_history(tr.t('hist.settings_applied'))

            if theme_changed:
                self._apply_theme(new_theme)
                self.theme_var.set(new_theme)
                self._populate_theme_menu()

            if lang_changed:
                self._switch_lang(new_lang)

        ttk.Button(win, text=tr.t('settings.save'),
                    command=save).pack(pady=15)

    # ========================================================================
    # КОНСОЛЬ
    # ========================================================================
    def _open_console(self):
        if (self._console_window is not None
                and self._console_window.winfo_exists()):
            self._console_window.lift()
            self._console_window.focus_force()
            return
        win = tk.Toplevel(self.root)
        win.title(tr.t('menu.file.console'))
        win.geometry("720x500")
        win.transient(self.root)
        self._console_window = win
        tk.Label(win,
                  text="Доступно: app, np, plt, tk, mpl, messagebox, tr.",
                  bg="#2b2b2b", fg="#cccccc", anchor=tk.W,
                  font=('Consolas', 9), padx=6, pady=4).pack(fill=tk.X)
        frame = ttk.Frame(win); frame.pack(fill=tk.BOTH, expand=True)
        output = tk.Text(frame, bg="#1e1e1e", fg="#d4d4d4",
                          insertbackground="#ffffff",
                          font=('Consolas', 10), wrap=tk.NONE, undo=True)
        output.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        sb_y = ttk.Scrollbar(frame, orient=tk.VERTICAL, command=output.yview)
        sb_y.pack(side=tk.RIGHT, fill=tk.Y)
        output.configure(yscrollcommand=sb_y.set)
        sb_x = ttk.Scrollbar(win, orient=tk.HORIZONTAL, command=output.xview)
        sb_x.pack(fill=tk.X)
        output.configure(xscrollcommand=sb_x.set)
        input_frame = ttk.Frame(win); input_frame.pack(fill=tk.X, padx=4, pady=4)
        ttk.Label(input_frame, text=">>>").pack(side=tk.LEFT, padx=(0, 4))
        entry = ttk.Entry(input_frame, font=('Consolas', 10))
        entry.pack(side=tk.LEFT, fill=tk.X, expand=True)
        entry.focus()
        namespace = {
            'app': self, 'np': np,
            'plt': __import__('matplotlib.pyplot', fromlist=['pyplot']),
            'tk': tk, 'mpl': mpl,
            'messagebox': messagebox, 'tr': tr,
            '__name__': '__console__', '__doc__': None,
        }
        multiline_buffer = {'text': None}
        history = []; history_pos = [-1]

        def write_output(text):
            output.insert(tk.END, text); output.see(tk.END)

        def execute(code_text):
            import io, contextlib
            buf = io.StringIO()
            try:
                with contextlib.redirect_stdout(buf):
                    try:
                        result = eval(code_text, namespace, namespace)
                        if result is not None:
                            print(repr(result))
                    except SyntaxError:
                        exec(code_text, namespace, namespace)
                out = buf.getvalue()
                if out:
                    write_output(out)
            except Exception as e:
                import traceback
                write_output("".join(traceback.format_exception(
                    type(e), e, e.__traceback__)) + "\n")

        def execute_with_history(code_text):
            if code_text.strip():
                history.append(code_text); history_pos[0] = -1
            execute(code_text)

        def on_enter(event=None):
            if multiline_buffer['text'] is not None:
                return
            code_text = entry.get()
            entry.delete(0, tk.END)
            write_output(f">>> {code_text}\n")
            if not code_text.strip():
                return
            execute_with_history(code_text)

        def on_ctrl_enter(event=None):
            if multiline_buffer['text'] is None:
                multiline_buffer['text'] = entry.get()
                entry.delete(0, tk.END)
                write_output("... " + multiline_buffer['text'] + "\n")
                return "break"
            else:
                multiline_buffer['text'] += "\n" + entry.get()
                entry.delete(0, tk.END)
                write_output("... " + entry.get() + "\n")
                execute_with_history(multiline_buffer['text'])
                multiline_buffer['text'] = None
                return "break"

        def on_up(event=None):
            if not history:
                return "break"
            if history_pos[0] == -1:
                history_pos[0] = len(history) - 1
            else:
                history_pos[0] = max(0, history_pos[0] - 1)
            entry.delete(0, tk.END); entry.insert(0, history[history_pos[0]])
            return "break"

        def on_down(event=None):
            if history_pos[0] == -1:
                return "break"
            history_pos[0] = min(len(history) - 1, history_pos[0] + 1)
            entry.delete(0, tk.END); entry.insert(0, history[history_pos[0]])
            return "break"

        entry.bind('<Return>', on_enter)
        entry.bind('<Control-Return>', on_ctrl_enter)
        entry.bind('<Up>', on_up)
        entry.bind('<Down>', on_down)
        write_output("GraphEditor Pro — Python-консоль\n\n")
        btn_frame = ttk.Frame(win); btn_frame.pack(fill=tk.X, padx=4, pady=(0, 6))

        def clear_output():
            output.delete("1.0", tk.END)

        def run_script():
            path = filedialog.askopenfilename(
                filetypes=[("Python", "*.py"), ("All files", "*.*")])
            if not path:
                return
            try:
                with open(path, 'r', encoding='utf-8') as f:
                    code = f.read()
                write_output(f"--- {os.path.basename(path)} ---\n")
                execute_with_history(code)
                write_output("--- Выполнено ---\n")
            except Exception as e:
                write_output(f"Ошибка: {e}\n")

        ttk.Button(btn_frame, text="Очистить", command=clear_output).pack(side=tk.LEFT, padx=2)
        ttk.Button(btn_frame, text="Загрузить .py...", command=run_script).pack(side=tk.LEFT, padx=2)
        ttk.Button(btn_frame, text="Закрыть", command=win.destroy).pack(side=tk.RIGHT, padx=2)
        win.after(100, entry.focus_set)

    # ========================================================================
    # ОНБОРДИНГ
    # ========================================================================
    def _maybe_start_onboarding(self):
        if not self.settings.get('show_tour_on_start', False):
            return
        if self.graphs_data:
            return
        self._start_onboarding()

    def _start_onboarding(self):
        self._onboarding_step = 0
        self._onboarding_widgets = [
            (self.notebook, 'tour.step1_title', 'tour.step1_text'),
            (self.graph_tree, 'tour.step2_title', 'tour.step2_text'),
            (self.plot_frame, 'tour.step3_title', 'tour.step3_text'),
            (self.anim_labelframe, 'tour.step4_title', 'tour.step4_text'),
            (self.plot_frame, 'tour.step5_title', 'tour.step5_text'),
        ]
        self._show_onboarding_step()

    def _show_onboarding_step(self):
        if self._onboarding_step >= len(self._onboarding_widgets):
            self._finish_onboarding(); return
        widget, title_key, text_key = self._onboarding_widgets[self._onboarding_step]
        try:
            widget.update_idletasks()
            x, y = widget.winfo_rootx(), widget.winfo_rooty()
            w, h = widget.winfo_width(), widget.winfo_height()
        except Exception:
            x, y, w, h = 300, 300, 200, 100
        sw, sh = self.root.winfo_screenwidth(), self.root.winfo_screenheight()
        popup = tk.Toplevel(self.root)
        popup.title(tr.t('msg.tour_title'))
        popup.transient(self.root)
        popup_w, popup_h = 340, 200
        px, py = x + w + 15, y
        if px + popup_w > sw - 20:
            px = x - popup_w - 15
            if px < 20:
                px = max(20, (sw - popup_w) // 2)
                py = min(y + h + 15, sh - popup_h - 40)
        py = max(20, min(py, sh - popup_h - 40))
        popup.geometry(f"{popup_w}x{popup_h}+{px}+{py}")
        total = len(self._onboarding_widgets)
        tk.Label(popup, text=tr.t('msg.tour_step',
                                    n=self._onboarding_step + 1, total=total),
                  font=('Arial', 8), fg='gray').pack(anchor=tk.W, padx=10, pady=(8, 0))
        tk.Label(popup, text=tr.t(title_key),
                  font=('Arial', 12, 'bold')).pack(anchor=tk.W, padx=10)
        tk.Label(popup, text=tr.t(text_key), wraplength=popup_w - 30,
                  justify=tk.LEFT).pack(anchor=tk.W, padx=10, pady=5, fill=tk.X)
        bf = tk.Frame(popup); bf.pack(side=tk.BOTTOM, fill=tk.X, padx=10, pady=8)

        def close_popup():
            try: popup.destroy()
            except Exception: pass
        def next_step():
            close_popup()
            self._onboarding_step += 1
            self.root.after(50, self._show_onboarding_step)
        def skip():
            close_popup(); self._finish_onboarding()

        tk.Button(bf, text=tr.t('msg.tour_skip'), command=skip).pack(side=tk.LEFT)
        is_last = self._onboarding_step == total - 1
        tk.Button(bf,
                   text=tr.t('msg.tour_finish') if is_last else tr.t('msg.tour_next'),
                   command=next_step, default=tk.ACTIVE).pack(side=tk.RIGHT)
        popup.bind('<Escape>', lambda e: skip())
        popup.protocol("WM_DELETE_WINDOW", skip)
        popup.after(50, lambda: popup.focus_force())

    def _finish_onboarding(self):
        self._show_history(tr.t('status.tour_finished'))